from __future__ import annotations

import asyncio
import base64
import hashlib
import io
import logging
import re
from time import monotonic
from typing import TypeVar

from telegram import (
    BotCommand,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    KeyboardButton,
    LinkPreviewOptions,
    ReplyKeyboardMarkup,
    Update,
)
from telegram.constants import ChatAction, ParseMode
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    filters,
)

from src.core.context_store import ContextStore, ConversationTurn
from src.core.model_preferences_store import ModelPreferencesStore
from src.core.rate_limiter import SlidingWindowRateLimiter
from src.core.user_assets_store import UserAsset, UserAssetsStore
from src.i18n import I18nService
from src.services.ollama_client import (
    OllamaClient,
    OllamaConnectionError,
    OllamaError,
    OllamaTimeoutError,
    WebModelInfo,
)
from src.services.model_orchestrator import ModelOrchestrator
from src.utils.telegram import split_message

logger = logging.getLogger(__name__)

MODEL_CALLBACK_PREFIX = "model:"
WEB_MODEL_CALLBACK_PREFIX = "webmodel:"
FILE_CALLBACK_PREFIX = "file:"
CLEAR_CALLBACK_PREFIX = "clear:"
MODEL_REFRESH_ACTION = "__refresh__"
MODEL_DEFAULT_ACTION = "__default__"
MODEL_PAGE_ACTION_PREFIX = "__page__:"
MODEL_CLOSE_ACTION = "__close__"
WEB_MODEL_REFRESH_ACTION = "__refresh__"
WEB_MODEL_PAGE_ACTION_PREFIX = "__page__:"
WEB_MODEL_CLOSE_ACTION = "__close__"
WEB_MODEL_DETAIL_ACTION = "__detail__:"
WEB_MODEL_DOWNLOAD_ACTION = "__download__:"
WEB_MODEL_SIZE_ACTION = "__size__:"
WEB_MODEL_CANCEL_ACTION = "__cancel__:"
WEB_MODEL_SEARCH_ACTION = "__search__"
DELETE_MODEL_CALLBACK_PREFIX = "delmod:"
DELETE_MODEL_CONFIRM_ACTION = "confirm"
DELETE_MODEL_ABORT_ACTION = "abort"
_WEB_MODELS_CACHE_TTL = 300.0  # 5 minutes
FILE_PAGE_ACTION = "page"
FILE_TOGGLE_ACTION = "toggle"
FILE_DELETE_ACTION = "delete"
FILE_ASK_ACTION = "ask"
FILE_CLOSE_ACTION = "close"
FILE_UPLOAD_ACTION = "upload"
FILE_PREVIEW_ACTION = "preview"
ICON_INFO = "ℹ️"
ICON_SUCCESS = "✅"
ICON_WARNING = "⚠️"
ICON_ERROR = "❌"
MODELS_PAGE_SIZE = 8
WEB_MODELS_PAGE_SIZE = 8
FILES_PAGE_SIZE = 6
FILES_CONTEXT_MAX_ITEMS_DEFAULT = 3
WEBSEARCH_MAX_RESULTS = 5
WEBSEARCH_CONTEXT_MAX_CHARS = 4000
FILES_CONTEXT_MAX_CHARS_DEFAULT = 6000


class BotHandlers:
    def __init__(
        self,
        ollama_client: OllamaClient,
        context_store: ContextStore,
        model_preferences_store: ModelPreferencesStore,
        user_assets_store: UserAssetsStore,
        default_model: str,
        use_chat_api: bool,
        keep_alive: str,
        image_max_bytes: int,
        document_max_bytes: int,
        document_max_chars: int,
        i18n: I18nService,
        files_context_max_items: int = FILES_CONTEXT_MAX_ITEMS_DEFAULT,
        files_context_max_chars: int = FILES_CONTEXT_MAX_CHARS_DEFAULT,
        allowed_user_ids: set[int] | None = None,
        rate_limiter: SlidingWindowRateLimiter | None = None,
    ) -> None:
        self._ollama_client = ollama_client
        self._context_store = context_store
        self._model_preferences_store = model_preferences_store
        self._user_assets_store = user_assets_store
        self._default_model = default_model
        self._use_chat_api = use_chat_api
        self._keep_alive = keep_alive
        self._image_max_bytes = image_max_bytes
        self._document_max_bytes = document_max_bytes
        self._document_max_chars = document_max_chars
        self._files_context_max_items = files_context_max_items
        self._files_context_max_chars = files_context_max_chars
        self._i18n = i18n
        self._allowed_user_ids = allowed_user_ids or set()
        self._rate_limiter = rate_limiter
        self._model_orchestrator = ModelOrchestrator(ollama_client)
        self._quick_action_map = self._build_quick_action_map()
        self._model_search_query_by_user: dict[int, str] = {}
        self._web_model_search_query_by_user: dict[int, str] = {}
        self._model_downloads_in_progress: set[str] = set()
        self._download_cancel_events: dict[str, asyncio.Event] = {}
        self._web_model_search_mode_users: set[int] = set()
        self._web_model_token_to_name: dict[str, str] = {}
        self._web_model_name_to_token: dict[str, str] = {}
        self._web_models_cache: list[WebModelInfo] = []
        self._web_models_cache_expires: float = 0.0
        self._askfile_target_by_user: dict[int, int] = {}
        self._upload_mode_users: set[int] = set()

    @staticmethod
    def required_i18n_keys() -> tuple[str, ...]:
        return (
            "commands.start",
            "commands.help",
            "commands.health",
            "commands.clear",
            "commands.models",
            "commands.webmodels",
            "commands.files",
            "commands.askfile",
            "commands.cancel",
            "commands.currentmodel",
            "ui.buttons.models",
            "ui.buttons.web_models",
            "ui.buttons.files",
            "ui.buttons.current_model",
            "ui.buttons.clear",
            "ui.buttons.help",
            "ui.buttons.open_models",
            "ui.buttons.use_default",
            "ui.buttons.refresh",
            "ui.buttons.refresh_models",
            "ui.buttons.prev_page",
            "ui.buttons.next_page",
            "ui.buttons.confirm",
            "ui.buttons.cancel",
            "ui.buttons.close",
            "ui.buttons.ask_file",
            "ui.input_placeholder",
            "messages.start_welcome",
            "messages.help",
            "messages.please_send_non_empty",
            "messages.askfile_usage",
            "messages.askfile_prompt",
            "messages.cancel_ask_done",
            "messages.cancel_nothing",
            "messages.voice_disabled",
            "messages.clear_confirm",
            "messages.clear_cancelled",
            "messages.clear_done",
            "messages.current_model",
            "messages.access_denied",
            "messages.access_denied_alert",
            "messages.rate_limit_exceeded",
            "messages.unexpected_error",
            "health.result",
            "health.ok",
            "health.degraded",
            "health.sqlite",
            "health.ollama",
            "health.ollama_ok_with_models",
            "health.runtime_ok",
            "health.latency",
            "models.available_title",
            "models.current_marker",
            "models.select_with",
            "models.tap_button",
            "models.no_matches",
            "models.page_status",
            "models.updated",
            "models.reset_default",
            "models.not_found",
            "models.not_available_anymore",
            "models.no_models_available",
            "files.available_title",
            "files.empty",
            "files.page_status",
            "files.instructions",
            "files.deleted",
            "files.not_found",
            "web_models.available_title",
            "web_models.select_with",
            "web_models.install_hint",
            "web_models.page_status",
            "web_models.no_matches",
            "web_models.no_models_available",
            "image.default_prompt",
            "image.model_without_vision",
            "image.too_large",
            "image.invalid_file",
            "image.processing_error",
            "image.read_error",
            "document.added",
            "document.too_large",
            "document.unsupported",
            "document.empty",
            "document.processing_error",
            "errors.ollama_timeout",
            "errors.ollama_connection",
            "errors.ollama_list_models",
            "errors.ollama_list_web_models",
            "errors.ollama_validate_model",
            "errors.save_model_preference",
            "errors.save_default_model_preference",
            "errors.files_storage",
            "errors.ollama_generic",
            "agent.planner_instruction",
            "agent.analyst_instruction",
            "agent.chat_instruction",
            "ui.buttons.add_file",
            "ui.buttons.preview",
            "files.upload_prompt",
            "files.upload_done",
            "orchestrator.switched_model",
            "orchestrator.task_vision",
            "orchestrator.task_code",
            "web_models.detail_title",
            "web_models.download_started",
            "web_models.download_done",
            "web_models.download_failed",
            "web_models.already_downloading",
            "web_models.size_select",
            "ui.buttons.download",
            "ui.buttons.open_web",
            "ui.buttons.search",
            "commands.deletemodel",
            "commands.info",
            "commands.websearch",
            "web_models.search_prompt",
            "web_models.download_cancelled",
            "models.delete_usage",
            "models.delete_confirm",
            "models.delete_done",
            "models.delete_failed",
            "models.delete_not_found",
            "models.info_not_found",
            "web_search.no_api_key",
            "web_search.searching",
            "web_search.no_results",
            "web_search.header",
            "web_search.sources_header",
            "web_search.usage",
        )

    async def set_commands(self, application: Application) -> None:
        for locale in self._i18n.available_locales:
            commands = [
                BotCommand(command="start", description=self._i18n.t("commands.start", locale=locale)),
                BotCommand(command="help", description=self._i18n.t("commands.help", locale=locale)),
                BotCommand(command="health", description=self._i18n.t("commands.health", locale=locale)),
                BotCommand(command="clear", description=self._i18n.t("commands.clear", locale=locale)),
                BotCommand(command="models", description=self._i18n.t("commands.models", locale=locale)),
                BotCommand(
                    command="webmodels",
                    description=self._i18n.t("commands.webmodels", locale=locale),
                ),
                BotCommand(command="files", description=self._i18n.t("commands.files", locale=locale)),
                BotCommand(command="askfile", description=self._i18n.t("commands.askfile", locale=locale)),
                BotCommand(command="cancel", description=self._i18n.t("commands.cancel", locale=locale)),
                BotCommand(
                    command="currentmodel",
                    description=self._i18n.t("commands.currentmodel", locale=locale),
                ),
                BotCommand(
                    command="deletemodel",
                    description=self._i18n.t("commands.deletemodel", locale=locale),
                ),
                BotCommand(
                    command="info",
                    description=self._i18n.t("commands.info", locale=locale),
                ),
                BotCommand(
                    command="websearch",
                    description=self._i18n.t("commands.websearch", locale=locale),
                ),
            ]
            if locale == self._i18n.default_locale:
                await application.bot.set_my_commands(commands)
            else:
                await application.bot.set_my_commands(commands, language_code=locale)

    async def start(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not await self._guard_access(update):
            return
        if not update.effective_message:
            return
        locale = self._locale(update)
        self._log_user_event("command_start", update)
        await update.effective_message.reply_text(
            self._i18n.t("messages.start_welcome", locale=locale),
            parse_mode=ParseMode.HTML,
            reply_markup=self._main_keyboard(locale),
        )

    async def help(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not await self._guard_access(update):
            return
        if not update.effective_message:
            return
        locale = self._locale(update)
        self._log_user_event("command_help", update)
        await update.effective_message.reply_text(
            self._i18n.t("messages.help", locale=locale),
            reply_markup=self._main_keyboard(locale),
        )

    async def health(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not await self._guard_access(update):
            return
        if not update.effective_message or not update.effective_user:
            return
        locale = self._locale(update)

        self._log_user_event("command_health", update)
        started_at = monotonic()

        db_ok = True
        db_detail = "OK"
        try:
            self._model_preferences_store.healthcheck()
        except Exception as error:
            logger.exception("health_db_check_failed user_id=%s", update.effective_user.id)
            db_ok = False
            db_detail = str(error)

        ollama_ok = True
        ollama_detail = "OK"
        ollama_model_count = 0
        try:
            models = await self._ollama_client.list_models()
            ollama_model_count = len(models)
            ollama_detail = self._i18n.t(
                "health.ollama_ok_with_models",
                locale=locale,
                count=ollama_model_count,
            )
        except Exception as error:
            logger.exception("health_ollama_check_failed user_id=%s", update.effective_user.id)
            ollama_ok = False
            ollama_detail = str(error)

        elapsed_ms = int((monotonic() - started_at) * 1000)
        overall_ok = db_ok and ollama_ok
        overall_text = (
            self._i18n.t("health.ok", locale=locale)
            if overall_ok
            else self._i18n.t("health.degraded", locale=locale)
        )

        lines = [self._info(self._i18n.t("health.result", locale=locale, status=overall_text))]
        lines.append(
            f"{ICON_SUCCESS if db_ok else ICON_ERROR} "
            f"{self._i18n.t('health.sqlite', locale=locale, detail=db_detail)}"
        )
        lines.append(
            f"{ICON_SUCCESS if ollama_ok else ICON_ERROR} "
            f"{self._i18n.t('health.ollama', locale=locale, detail=ollama_detail)}"
        )
        lines.append(f"{ICON_INFO} {self._i18n.t('health.runtime_ok', locale=locale)}")
        lines.append(f"{ICON_INFO} {self._i18n.t('health.latency', locale=locale, ms=elapsed_ms)}")

        logger.info(
            "healthcheck_result user_id=%s overall_ok=%s db_ok=%s ollama_ok=%s ollama_models=%d elapsed_ms=%d",
            update.effective_user.id,
            overall_ok,
            db_ok,
            ollama_ok,
            ollama_model_count,
            elapsed_ms,
        )

        await update.effective_message.reply_text(
            "\n".join(lines), reply_markup=self._main_keyboard(locale)
        )

    async def clear(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not await self._guard_access(update):
            return
        if not update.effective_message or not update.effective_user:
            return
        locale = self._locale(update)
        self._log_user_event("command_clear", update)
        await update.effective_message.reply_text(
            self._warning(self._i18n.t("messages.clear_confirm", locale=locale)),
            parse_mode=ParseMode.HTML,
            reply_markup=self._clear_inline_keyboard(locale),
        )

    async def web_models(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not await self._guard_access(update):
            return
        if not update.effective_message or not update.effective_user:
            return

        locale = self._locale(update)
        user_id = update.effective_user.id
        self._log_user_event("command_webmodels", update)
        search_query = " ".join(context.args).strip() if context.args else ""

        try:
            models = await self._fetch_web_models()
        except OllamaTimeoutError:
            await update.effective_message.reply_text(
                self._warning(self._i18n.t("errors.ollama_timeout", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaConnectionError:
            await update.effective_message.reply_text(
                self._error(self._i18n.t("errors.ollama_connection", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaError as error:
            logger.warning("Ollama error while listing web models: %s", error)
            await update.effective_message.reply_text(
                self._error(self._i18n.t("errors.ollama_list_web_models", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        if not models:
            await update.effective_message.reply_text(
                self._info(self._i18n.t("web_models.no_models_available", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        self._web_model_search_query_by_user[user_id] = search_query
        filtered_models = self._filter_web_models(models, search_query)
        if not filtered_models:
            await update.effective_message.reply_text(
                self._warning(self._i18n.t("web_models.no_matches", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        page = 1
        page_models, total_pages = self._paginate_items(filtered_models, page, WEB_MODELS_PAGE_SIZE)

        lines = [self._info(self._i18n.t("web_models.available_title", locale=locale))]
        for m in page_models:
            badges = []
            if "vision" in m.capabilities:
                badges.append("👁")
            if "thinking" in m.capabilities:
                badges.append("💭")
            line = f"- {m.name}"
            if badges:
                line += "  " + " ".join(badges)
            if m.sizes:
                line += "  📦 " + " · ".join(m.sizes[:4])
            lines.append(line)
        lines.append("")
        lines.append(self._i18n.t("web_models.page_status", locale=locale, page=page, pages=total_pages))
        lines.append(self._i18n.t("web_models.select_with", locale=locale))

        await update.effective_message.reply_text(
            "\n".join(lines),
            reply_markup=self._web_models_inline_keyboard(
                locale=locale,
                models=page_models,
                page=page,
                total_pages=total_pages,
            ),
        )

    async def files(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not await self._guard_access(update):
            return
        if not update.effective_message or not update.effective_user:
            return

        locale = self._locale(update)
        user_id = update.effective_user.id
        self._log_user_event("command_files", update)

        try:
            assets = self._user_assets_store.list_assets(user_id)
        except Exception as error:
            logger.exception("Failed to list user assets: %s", error)
            await update.effective_message.reply_text(
                self._error(self._i18n.t("errors.files_storage", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        if not assets:
            await update.effective_message.reply_text(
                self._info(self._i18n.t("files.empty", locale=locale)),
                reply_markup=InlineKeyboardMarkup(
                    [
                        [
                            InlineKeyboardButton(
                                text=self._i18n.t("ui.buttons.add_file", locale=locale),
                                callback_data=f"{FILE_CALLBACK_PREFIX}{FILE_UPLOAD_ACTION}",
                            )
                        ]
                    ]
                ),
            )
            return

        page = 1
        page_assets, total_pages = self._paginate_assets(assets, page)
        text = self._files_page_text(locale=locale, assets=page_assets, page=page, total_pages=total_pages)
        await update.effective_message.reply_text(
            text,
            reply_markup=self._files_inline_keyboard(
                locale=locale,
                assets=page_assets,
                page=page,
                total_pages=total_pages,
            ),
        )

    async def askfile(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not await self._guard_access(update, apply_rate_limit=True):
            return
        if not update.effective_message or not update.effective_user:
            return

        locale = self._locale(update)
        user_id = update.effective_user.id
        self._log_user_event("command_askfile", update)

        if not context.args or len(context.args) < 2:
            await update.effective_message.reply_text(
                self._warning(self._i18n.t("messages.askfile_usage", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        asset_id_raw = context.args[0].strip()
        prompt = " ".join(context.args[1:]).strip()
        if not asset_id_raw.isdigit() or not prompt:
            await update.effective_message.reply_text(
                self._warning(self._i18n.t("messages.askfile_usage", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        asset_id = int(asset_id_raw)
        try:
            asset = self._user_assets_store.get_asset(user_id, asset_id)
        except Exception as error:
            logger.exception("Failed to read askfile asset: %s", error)
            await update.effective_message.reply_text(
                self._error(self._i18n.t("errors.files_storage", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        if not asset:
            await update.effective_message.reply_text(
                self._warning(self._i18n.t("files.not_found", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        await self._answer_with_asset(
            update=update,
            locale=locale,
            user_id=user_id,
            asset=asset,
            prompt=prompt,
        )

    async def _answer_with_asset(
        self,
        *,
        update: Update,
        locale: str,
        user_id: int,
        asset: UserAsset,
        prompt: str,
    ) -> None:
        turns = self._context_store.get_turns(user_id)
        model = self._get_user_model(user_id)
        started_at = monotonic()
        agent_name = self._select_agent(prompt)
        system_instruction = self._agent_system_instruction(agent_name, locale)

        if asset.asset_kind == "image":
            # Pass the stored image bytes on the current user message so vision
            # models actually see the pixels.  The text analysis is added as a
            # prior assistant turn so text-only models still have useful context.
            extra_turns = self._build_image_context_turns([asset])
            prompt_to_send = prompt
            prompt_images: list[str] | None = [asset.image_base64] if asset.image_base64 else None
        else:
            extra_turns = None
            prompt_to_send = self._augment_prompt_with_assets(
                prompt=prompt, assets=[asset], force_single=True
            )
            prompt_images = None

        # Orchestrate model selection for this asset type
        model, orch_notification, vision_found = await self._orchestrate_model(
            prompt=prompt_to_send,
            has_images=bool(prompt_images),
            preferred_model=model,
            locale=locale,
        )
        if prompt_images and not vision_found:
            prompt_images = None
            orch_notification = None
            logger.warning(
                "answer_with_asset no_vision_model asset_id=%s using_text_descriptions",
                asset.id,
            )

        await update.effective_chat.send_action(action=ChatAction.TYPING)
        try:
            ollama_response = await self._generate_response(
                user_id=user_id,
                model=model,
                prompt=prompt_to_send,
                turns=turns,
                system_instruction=system_instruction,
                extra_turns=extra_turns or None,
                prompt_images=prompt_images,
            )
        except OllamaTimeoutError:
            await update.effective_message.reply_text(
                self._warning(self._i18n.t("errors.ollama_timeout", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaConnectionError:
            await update.effective_message.reply_text(
                self._error(self._i18n.t("errors.ollama_connection", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaError as error:
            logger.warning("Ollama askfile error: %s", error)
            await update.effective_message.reply_text(
                self._error(self._i18n.t("errors.ollama_generic", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        self._context_store.append(user_id, role="user", content=f"[AskFile #{asset.id}] {prompt}")
        self._context_store.append(user_id, role="assistant", content=ollama_response.text)

        elapsed_ms = int((monotonic() - started_at) * 1000)
        logger.info(
            "askfile_completed user_id=%s model=%s file_id=%d agent=%s input_chars=%d output_chars=%d elapsed_ms=%d",
            user_id,
            model,
            asset.id,
            agent_name,
            len(prompt),
            len(ollama_response.text),
            elapsed_ms,
        )

        for chunk in split_message(ollama_response.text):
            await update.effective_message.reply_text(chunk, parse_mode=ParseMode.HTML)
        if orch_notification:
            await update.effective_message.reply_text(
                f"<i>{orch_notification}</i>", parse_mode=ParseMode.HTML
            )

    async def cancel(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        """Cancel any pending interaction mode (e.g. inline ask)."""
        if not await self._guard_access(update):
            return
        if not update.effective_message or not update.effective_user:
            return
        locale = self._locale(update)
        user_id = update.effective_user.id
        was_upload = user_id in self._upload_mode_users
        self._upload_mode_users.discard(user_id)
        was_web_search_mode = user_id in self._web_model_search_mode_users
        self._web_model_search_mode_users.discard(user_id)
        pending = self._askfile_target_by_user.pop(user_id, None)
        if pending is not None or was_upload or was_web_search_mode:
            await update.effective_message.reply_text(
                self._i18n.t("messages.cancel_ask_done", locale=locale),
                reply_markup=self._main_keyboard(locale),
            )
        else:
            await update.effective_message.reply_text(
                self._i18n.t("messages.cancel_nothing", locale=locale),
                reply_markup=self._main_keyboard(locale),
            )

    async def models(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not await self._guard_access(update):
            return
        if not update.effective_message or not update.effective_user:
            return

        locale = self._locale(update)
        user_id = update.effective_user.id
        self._log_user_event("command_models", update)
        requested_model = " ".join(context.args).strip() if context.args else ""

        try:
            models = await self._ollama_client.list_models()
        except OllamaTimeoutError:
            await update.effective_message.reply_text(
                self._warning(self._i18n.t("errors.ollama_timeout", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaConnectionError:
            await update.effective_message.reply_text(
                self._error(self._i18n.t("errors.ollama_connection", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaError as error:
            logger.warning("Ollama error while listing models: %s", error)
            await update.effective_message.reply_text(
                self._error(self._i18n.t("errors.ollama_list_models", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        if not models:
            await update.effective_message.reply_text(
                self._info(self._i18n.t("models.no_models_available", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        if requested_model and (
            requested_model in models or self._ollama_client.can_use_cloud_model(requested_model)
        ):
            try:
                self._model_preferences_store.set_user_model(user_id, requested_model)
            except Exception as error:
                logger.exception("Failed to save user model preference: %s", error)
                await update.effective_message.reply_text(
                    self._error(self._i18n.t("errors.save_model_preference", locale=locale)),
                    reply_markup=self._main_keyboard(locale),
                )
                return
            await update.effective_message.reply_text(
                self._success(self._i18n.t("models.updated", locale=locale, model=requested_model)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        current_model = self._get_user_model(user_id)
        search_query = requested_model if requested_model else ""
        self._model_search_query_by_user[user_id] = search_query

        filtered_models = self._filter_models(models, search_query)
        if not filtered_models:
            await update.effective_message.reply_text(
                self._warning(self._i18n.t("models.no_matches", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        page = 1
        page_models, total_pages = self._paginate_models(filtered_models, page)

        lines = [self._info(self._i18n.t("models.available_title", locale=locale))]
        for model in page_models:
            marker = self._i18n.t("models.current_marker", locale=locale) if model == current_model else ""
            lines.append(f"- {model}{marker}")
        lines.append("")
        lines.append(self._i18n.t("models.page_status", locale=locale, page=page, pages=total_pages))
        lines.append(self._i18n.t("models.select_with", locale=locale))
        lines.append(self._i18n.t("models.tap_button", locale=locale))

        inline_keyboard = self._models_inline_keyboard(
            locale,
            models=page_models,
            current_model=current_model,
            page=page,
            total_pages=total_pages,
        )
        await update.effective_message.reply_text(
            "\n".join(lines),
            reply_markup=inline_keyboard,
        )

    async def clear_callback(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not await self._guard_access(update):
            return
        query = update.callback_query
        if not query or not update.effective_user:
            return
        if not query.message:
            return
        locale = self._locale(update)

        await query.answer()

        data = query.data or ""
        if not data.startswith(CLEAR_CALLBACK_PREFIX):
            return

        action = data.removeprefix(CLEAR_CALLBACK_PREFIX).strip()
        if action == "cancel":
            await query.message.reply_text(
                self._info(self._i18n.t("messages.clear_cancelled", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        if action == "confirm":
            self._context_store.clear(update.effective_user.id)
            await query.message.reply_text(
                self._success(self._i18n.t("messages.clear_done", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )

    async def select_model_callback(
        self, update: Update, context: ContextTypes.DEFAULT_TYPE
    ) -> None:
        if not await self._guard_access(update):
            return
        query = update.callback_query
        if not query or not update.effective_user:
            return
        if not query.message:
            return
        locale = self._locale(update)

        await query.answer()

        data = query.data or ""
        if not data.startswith(MODEL_CALLBACK_PREFIX):
            return

        selected_model = data.removeprefix(MODEL_CALLBACK_PREFIX).strip()
        if not selected_model:
            return

        if selected_model == MODEL_REFRESH_ACTION:
            self._model_search_query_by_user[update.effective_user.id] = ""
            await self._show_models_page(update, 1)
            return

        if selected_model == MODEL_CLOSE_ACTION:
            self._model_search_query_by_user.pop(update.effective_user.id, None)
            try:
                await query.message.delete()
            except Exception as error:
                logger.debug("Failed to delete local models message: %s", error)
                await query.edit_message_reply_markup(reply_markup=None)
            return

        if selected_model.startswith(MODEL_PAGE_ACTION_PREFIX):
            page_raw = selected_model.removeprefix(MODEL_PAGE_ACTION_PREFIX).strip()
            if not page_raw.isdigit():
                return
            await self._show_models_page(update, int(page_raw))
            return

        user_id = update.effective_user.id

        try:
            models = await self._ollama_client.list_models()
        except OllamaTimeoutError:
            await query.message.reply_text(
                self._warning(self._i18n.t("errors.ollama_timeout", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaConnectionError:
            await query.message.reply_text(
                self._error(self._i18n.t("errors.ollama_connection", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaError as error:
            logger.warning("Ollama error while selecting model: %s", error)
            await query.message.reply_text(
                self._error(self._i18n.t("errors.ollama_validate_model", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        if selected_model not in models:
            if selected_model == MODEL_DEFAULT_ACTION:
                try:
                    self._model_preferences_store.set_user_model(user_id, self._default_model)
                except Exception as error:
                    logger.exception("Failed to save default model preference: %s", error)
                    await query.message.reply_text(
                        self._error(
                            self._i18n.t("errors.save_default_model_preference", locale=locale)
                        ),
                        reply_markup=self._main_keyboard(locale),
                    )
                    return

                await query.message.reply_text(
                    self._success(
                        self._i18n.t("models.reset_default", locale=locale, model=self._default_model)
                    ),
                    reply_markup=self._main_keyboard(locale),
                )
                return

            await query.message.reply_text(
                self._warning(self._i18n.t("models.not_available_anymore", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        try:
            self._model_preferences_store.set_user_model(user_id, selected_model)
        except Exception as error:
            logger.exception("Failed to save user model preference: %s", error)
            await query.message.reply_text(
                self._error(self._i18n.t("errors.save_model_preference", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        await query.message.reply_text(
            self._success(self._i18n.t("models.updated", locale=locale, model=selected_model)),
            reply_markup=self._main_keyboard(locale),
        )

    async def select_web_model_callback(
        self, update: Update, context: ContextTypes.DEFAULT_TYPE
    ) -> None:
        if not await self._guard_access(update):
            return
        query = update.callback_query
        if not query or not update.effective_user or not query.message:
            return

        locale = self._locale(update)
        await query.answer()

        data = query.data or ""
        if not data.startswith(WEB_MODEL_CALLBACK_PREFIX):
            return

        action = data.removeprefix(WEB_MODEL_CALLBACK_PREFIX).strip()
        if action == WEB_MODEL_SEARCH_ACTION:
            user_id = update.effective_user.id
            self._web_model_search_mode_users.add(user_id)
            await query.answer()
            await query.message.reply_text(
                self._i18n.t("web_models.search_prompt", locale=locale),
            )
            return

        if action == WEB_MODEL_REFRESH_ACTION:
            self._web_model_search_query_by_user[update.effective_user.id] = ""
            await self._show_web_models_page(update, 1, force_refresh=True)
            return

        if action == WEB_MODEL_CLOSE_ACTION:
            self._web_model_search_query_by_user.pop(update.effective_user.id, None)
            try:
                await query.message.delete()
            except Exception as error:
                logger.debug("Failed to delete web models message: %s", error)
                await query.edit_message_reply_markup(reply_markup=None)
            return

        if action.startswith(WEB_MODEL_PAGE_ACTION_PREFIX):
            page_raw = action.removeprefix(WEB_MODEL_PAGE_ACTION_PREFIX).strip()
            if not page_raw.isdigit():
                return
            await self._show_web_models_page(update, int(page_raw))
            return

        if action.startswith(WEB_MODEL_CANCEL_ACTION):
            callback_model = action.removeprefix(WEB_MODEL_CANCEL_ACTION).strip()
            model_name = self._resolve_web_model_callback_value(callback_model)
            cancel_ev = self._download_cancel_events.get(model_name)
            if cancel_ev:
                cancel_ev.set()
            return

        if action.startswith(WEB_MODEL_DETAIL_ACTION):
            callback_model = action.removeprefix(WEB_MODEL_DETAIL_ACTION).strip()
            model_name = self._resolve_web_model_callback_value(callback_model)
            if not model_name:
                return
            try:
                all_models = await self._fetch_web_models()
            except (OllamaError, OllamaTimeoutError, OllamaConnectionError):
                all_models = []
            info = next((m for m in all_models if m.name == model_name), None)
            await self._edit_models_message(
                query=query,
                text=self._format_web_model_detail(info, model_name),
                reply_markup=InlineKeyboardMarkup(
                    [
                        [
                            InlineKeyboardButton(
                                text=self._i18n.t("ui.buttons.download", locale=locale),
                                callback_data=(
                                    f"{WEB_MODEL_CALLBACK_PREFIX}{WEB_MODEL_DOWNLOAD_ACTION}"
                                    f"{self._web_model_token(model_name)}"
                                ),
                            ),
                            InlineKeyboardButton(
                                text=self._i18n.t("ui.buttons.open_web", locale=locale),
                                url=f"https://ollama.com/library/{model_name}",
                            ),
                        ],
                        [
                            InlineKeyboardButton(
                                text=self._i18n.t("ui.buttons.close", locale=locale),
                                callback_data=f"{WEB_MODEL_CALLBACK_PREFIX}{WEB_MODEL_CLOSE_ACTION}",
                            ),
                        ],
                    ]
                ),
            )
            return

        if action.startswith(WEB_MODEL_DOWNLOAD_ACTION):
            callback_model = action.removeprefix(WEB_MODEL_DOWNLOAD_ACTION).strip()
            model_name = self._resolve_web_model_callback_value(callback_model)
            if not model_name:
                return
            # Look up sizes to offer sub-selection
            try:
                all_models = await self._fetch_web_models()
            except (OllamaError, OllamaTimeoutError, OllamaConnectionError):
                all_models = []
            info = next((m for m in all_models if m.name == model_name), None)
            if info and len(info.sizes) > 1:
                # Show size selection keyboard
                size_rows: list[list[InlineKeyboardButton]] = []
                size_row: list[InlineKeyboardButton] = []
                for size in info.sizes:
                    full = f"{model_name}:{size}"
                    callback_payload = f"{self._web_model_token(model_name)}:{size}"
                    size_row.append(
                        InlineKeyboardButton(
                            text=size,
                            callback_data=f"{WEB_MODEL_CALLBACK_PREFIX}{WEB_MODEL_SIZE_ACTION}{callback_payload}",
                        )
                    )
                    if len(size_row) == 3:
                        size_rows.append(size_row)
                        size_row = []
                if size_row:
                    size_rows.append(size_row)
                size_rows.append(
                    [
                        InlineKeyboardButton(
                            text=f"{self._i18n.t('ui.buttons.download', locale=locale)} (latest)",
                            callback_data=(
                                f"{WEB_MODEL_CALLBACK_PREFIX}{WEB_MODEL_SIZE_ACTION}"
                                f"{self._web_model_token(model_name)}:latest"
                            ),
                        )
                    ]
                )
                size_rows.append(
                    [
                        InlineKeyboardButton(
                            text=self._i18n.t("ui.buttons.close", locale=locale),
                            callback_data=f"{WEB_MODEL_CALLBACK_PREFIX}{WEB_MODEL_CLOSE_ACTION}",
                        )
                    ]
                )
                await self._edit_models_message(
                    query=query,
                    text=self._info(
                        self._i18n.t("web_models.size_select", locale=locale, model=model_name)
                    ),
                    reply_markup=InlineKeyboardMarkup(size_rows),
                )
                return
            # No size selection needed → download directly
            if model_name in self._model_downloads_in_progress:
                await query.answer(
                    self._i18n.t(
                        "web_models.already_downloading", locale=locale, model=model_name
                    ),
                    show_alert=True,
                )
                return
            self._model_downloads_in_progress.add(model_name)
            cancel_event = asyncio.Event()
            self._download_cancel_events[model_name] = cancel_event
            callback_model = self._web_model_token(model_name)
            cancel_kb = InlineKeyboardMarkup(
                [
                    [
                        InlineKeyboardButton(
                            text=self._i18n.t("ui.buttons.cancel", locale=locale),
                            callback_data=f"{WEB_MODEL_CALLBACK_PREFIX}{WEB_MODEL_CANCEL_ACTION}{callback_model}",
                        )
                    ]
                ]
            )
            try:
                await self._edit_models_message(
                    query=query,
                    text=self._info(
                        self._i18n.t(
                            "web_models.download_started", locale=locale, model=model_name
                        )
                    ),
                    reply_markup=cancel_kb,
                )
            except Exception as edit_error:  # noqa: BLE001
                logger.debug("Could not edit message for download start: %s", edit_error)
            asyncio.create_task(
                self._background_pull_model(
                    chat_id=query.message.chat_id,
                    message_id=query.message.message_id,
                    model_name=model_name,
                    locale=locale,
                    context=context,
                    cancel_event=cancel_event,
                )
            )
            return

        if action.startswith(WEB_MODEL_SIZE_ACTION):
            payload = action.removeprefix(WEB_MODEL_SIZE_ACTION).strip()
            # payload is "model_token:size_tag" — rpartition to split on last ":"
            callback_model, _, size_tag = payload.rpartition(":")
            model_name = self._resolve_web_model_callback_value(callback_model)
            if not model_name or not size_tag:
                return
            full_model = f"{model_name}:{size_tag}"
            if full_model in self._model_downloads_in_progress or model_name in self._model_downloads_in_progress:
                await query.answer(
                    self._i18n.t(
                        "web_models.already_downloading", locale=locale, model=full_model
                    ),
                    show_alert=True,
                )
                return
            self._model_downloads_in_progress.add(full_model)
            cancel_event = asyncio.Event()
            self._download_cancel_events[full_model] = cancel_event
            callback_model = self._web_model_token(full_model)
            cancel_kb = InlineKeyboardMarkup(
                [
                    [
                        InlineKeyboardButton(
                            text=self._i18n.t("ui.buttons.cancel", locale=locale),
                            callback_data=f"{WEB_MODEL_CALLBACK_PREFIX}{WEB_MODEL_CANCEL_ACTION}{callback_model}",
                        )
                    ]
                ]
            )
            try:
                await self._edit_models_message(
                    query=query,
                    text=self._info(
                        self._i18n.t(
                            "web_models.download_started", locale=locale, model=full_model
                        )
                    ),
                    reply_markup=cancel_kb,
                )
            except Exception as edit_error:  # noqa: BLE001
                logger.debug("Could not edit message for size download start: %s", edit_error)
            asyncio.create_task(
                self._background_pull_model(
                    chat_id=query.message.chat_id,
                    message_id=query.message.message_id,
                    model_name=full_model,
                    locale=locale,
                    context=context,
                    cancel_event=cancel_event,
                )
            )
            return

    async def _background_pull_model(
        self,
        *,
        chat_id: int,
        message_id: int,
        model_name: str,
        locale: str,
        context: ContextTypes.DEFAULT_TYPE,
        cancel_event: asyncio.Event,
    ) -> None:
        callback_model = self._web_model_token(model_name)
        cancel_kb = InlineKeyboardMarkup(
            [
                [
                    InlineKeyboardButton(
                        text=self._i18n.t("ui.buttons.cancel", locale=locale),
                        callback_data=f"{WEB_MODEL_CALLBACK_PREFIX}{WEB_MODEL_CANCEL_ACTION}{callback_model}",
                    )
                ]
            ]
        )
        last_edit: list[float] = [monotonic()]

        async def _on_progress(status: str, completed: int, total: int) -> None:
            now = monotonic()
            if total > 0 and completed > 0:
                pct = completed / total
                filled = int(pct * 10)
                bar = "█" * filled + "░" * (10 - filled)
                mb_done = completed / 1_048_576
                mb_total = total / 1_048_576
                progress_line = f"{bar} {pct * 100:.0f}%  ({mb_done:.1f}/{mb_total:.1f} MB)"
            else:
                progress_line = status or "…"
            if now - last_edit[0] < 2.0:
                return
            last_edit[0] = now
            try:
                await context.bot.edit_message_text(
                    chat_id=chat_id,
                    message_id=message_id,
                    text=f"⏳ <b>{model_name}</b>\n{progress_line}",
                    parse_mode=ParseMode.HTML,
                    reply_markup=cancel_kb,
                )
            except Exception:  # noqa: BLE001
                pass

        result_text: str | None = None
        try:
            await self._ollama_client.pull_model(
                model_name,
                progress_callback=_on_progress,
                cancel_event=cancel_event,
            )
            result_text = self._success(
                self._i18n.t("web_models.download_done", locale=locale, model=model_name)
            )
        except asyncio.CancelledError:
            result_text = self._i18n.t(
                "web_models.download_cancelled", locale=locale, model=model_name
            )
        except (OllamaError, OllamaTimeoutError, OllamaConnectionError) as error:
            logger.error("pull_model_failed model=%s error=%s", model_name, error)
            result_text = self._error(
                self._i18n.t(
                    "web_models.download_failed", locale=locale, model=model_name, error=str(error)
                )
            )
        except Exception as error:  # noqa: BLE001
            logger.exception("pull_model_unexpected_error model=%s error=%s", model_name, error)
            result_text = self._error(
                self._i18n.t(
                    "web_models.download_failed", locale=locale, model=model_name, error=str(error)
                )
            )
        finally:
            self._model_downloads_in_progress.discard(model_name)
            self._download_cancel_events.pop(model_name, None)

        if result_text:
            try:
                await context.bot.edit_message_text(
                    chat_id=chat_id,
                    message_id=message_id,
                    text=result_text,
                    reply_markup=None,
                )
            except Exception:  # noqa: BLE001
                try:
                    await context.bot.send_message(chat_id=chat_id, text=result_text)
                except Exception as notify_error:  # noqa: BLE001
                    logger.warning(
                        "pull_model_notify_failed chat_id=%s model=%s error=%s",
                        chat_id,
                        model_name,
                        notify_error,
                    )

    # ------------------------------------------------------------------ #
    #  /deletemodel                                                        #
    # ------------------------------------------------------------------ #

    async def delete_model(
        self, update: Update, context: ContextTypes.DEFAULT_TYPE
    ) -> None:
        if not await self._guard_access(update):
            return
        if not update.effective_message or not update.effective_user:
            return
        locale = self._locale(update)

        args = context.args or []
        if not args:
            await update.effective_message.reply_text(
                self._warning(self._i18n.t("models.delete_usage", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        model_name = args[0].strip()
        confirm_text = self._i18n.t("models.delete_confirm", locale=locale, model=model_name)
        keyboard = InlineKeyboardMarkup(
            [
                [
                    InlineKeyboardButton(
                        text=self._i18n.t("ui.buttons.confirm", locale=locale),
                        callback_data=f"{DELETE_MODEL_CALLBACK_PREFIX}{DELETE_MODEL_CONFIRM_ACTION}:{model_name}",
                    ),
                    InlineKeyboardButton(
                        text=self._i18n.t("ui.buttons.cancel", locale=locale),
                        callback_data=f"{DELETE_MODEL_CALLBACK_PREFIX}{DELETE_MODEL_ABORT_ACTION}",
                    ),
                ]
            ]
        )
        await update.effective_message.reply_text(
            confirm_text,
            reply_markup=keyboard,
        )

    async def delete_model_callback(
        self, update: Update, context: ContextTypes.DEFAULT_TYPE
    ) -> None:
        if not await self._guard_access(update):
            return
        query = update.callback_query
        if not query or not query.message:
            return
        await query.answer()

        locale = self._locale(update)
        data = query.data or ""
        payload = data.removeprefix(DELETE_MODEL_CALLBACK_PREFIX).strip()

        if payload == DELETE_MODEL_ABORT_ACTION:
            try:
                await query.message.delete()
            except Exception:  # noqa: BLE001
                await query.edit_message_reply_markup(reply_markup=None)
            return

        if payload.startswith(f"{DELETE_MODEL_CONFIRM_ACTION}:"):
            model_name = payload.removeprefix(f"{DELETE_MODEL_CONFIRM_ACTION}:").strip()
            try:
                await self._ollama_client.delete_model(model_name)
                text = self._success(
                    self._i18n.t("models.delete_done", locale=locale, model=model_name)
                )
            except OllamaError as error:
                err_str = str(error)
                if "not found" in err_str.lower():
                    text = self._warning(
                        self._i18n.t("models.delete_not_found", locale=locale, model=model_name)
                    )
                else:
                    text = self._error(
                        self._i18n.t(
                            "models.delete_failed",
                            locale=locale,
                            model=model_name,
                            error=err_str,
                        )
                    )
            except Exception as error:  # noqa: BLE001
                text = self._error(
                    self._i18n.t(
                        "models.delete_failed",
                        locale=locale,
                        model=model_name,
                        error=str(error),
                    )
                )
            try:
                await query.edit_message_text(text, reply_markup=None)
            except Exception:  # noqa: BLE001
                await query.message.reply_text(text)

    # ------------------------------------------------------------------ #
    #  /info                                                               #
    # ------------------------------------------------------------------ #

    async def model_info(
        self, update: Update, context: ContextTypes.DEFAULT_TYPE
    ) -> None:
        if not await self._guard_access(update):
            return
        if not update.effective_message or not update.effective_user:
            return
        locale = self._locale(update)
        user_id = update.effective_user.id

        args = context.args or []
        if args:
            model_name = args[0].strip()
        else:
            model_name = (
                self._model_preferences_store.get_user_model(user_id) or self._default_model
            )

        try:
            data = await self._ollama_client.show_model(model_name)
        except OllamaError as error:
            if "not found" in str(error).lower():
                await update.effective_message.reply_text(
                    self._warning(
                        self._i18n.t("models.info_not_found", locale=locale, model=model_name)
                    ),
                    reply_markup=self._main_keyboard(locale),
                )
            else:
                await update.effective_message.reply_text(
                    self._error(str(error)),
                    reply_markup=self._main_keyboard(locale),
                )
            return
        except (OllamaTimeoutError, OllamaConnectionError) as error:
            await update.effective_message.reply_text(
                self._error(str(error)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        # Build info card
        modelfile = data.get("modelfile", "")
        details = data.get("details", {})
        family = details.get("family", "—")
        param_size = details.get("parameter_size", "—")
        quant = details.get("quantization_level", "—")
        arch = details.get("architecture", details.get("family", "—"))
        size_bytes = data.get("size", 0)
        size_mb = size_bytes / 1_048_576 if size_bytes else 0

        # Extract system prompt if present
        system_lines = [
            line.removeprefix("SYSTEM ").strip()
            for line in modelfile.splitlines()
            if line.upper().startswith("SYSTEM ")
        ]
        system_prompt = system_lines[0][:200] if system_lines else None

        lines = [
            f"<b>{model_name}</b>",
            "",
            f"🏗 <b>Family:</b> {family}",
            f"⚙️ <b>Params:</b> {param_size}",
            f"🗜 <b>Quant:</b> {quant}",
            f"📐 <b>Arch:</b> {arch}",
        ]
        if size_mb:
            lines.append(f"💾 <b>Size:</b> {size_mb:.0f} MB")
        if system_prompt:
            lines.append(f"\n📝 <b>System:</b> {system_prompt}")

        await update.effective_message.reply_text(
            "\n".join(lines),
            parse_mode=ParseMode.HTML,
            reply_markup=self._main_keyboard(locale),
        )

    # ------------------------------------------------------------------ #
    #  /websearch                                                          #
    # ------------------------------------------------------------------ #

    async def web_search_cmd(
        self, update: Update, context: ContextTypes.DEFAULT_TYPE
    ) -> None:
        if not await self._guard_access(update, apply_rate_limit=True):
            return
        if not update.effective_message or not update.effective_user:
            return
        locale = self._locale(update)
        user_id = update.effective_user.id

        query_str = " ".join(context.args).strip() if context.args else ""
        if not query_str:
            await update.effective_message.reply_text(
                self._warning(self._i18n.t("web_search.usage", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        # Check API key
        if not self._ollama_client.web_search_available:
            await update.effective_message.reply_text(
                self._warning(self._i18n.t("web_search.no_api_key", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        # Inform user
        await update.effective_message.reply_text(
            self._info(self._i18n.t("web_search.searching", locale=locale, query=query_str)),
        )
        await update.effective_chat.send_action(action=ChatAction.TYPING)

        try:
            results = await self._ollama_client.web_search(
                query_str, max_results=WEBSEARCH_MAX_RESULTS
            )
        except (OllamaError, OllamaTimeoutError, OllamaConnectionError) as error:
            await update.effective_message.reply_text(
                self._error(str(error)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        if not results:
            await update.effective_message.reply_text(
                self._info(self._i18n.t("web_search.no_results", locale=locale, query=query_str)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        # Build context block from results
        context_parts: list[str] = [
            self._i18n.t("web_search.header", locale=locale, query=query_str)
        ]
        for i, r in enumerate(results, 1):
            title = r.get("title", "")
            url = r.get("url", "")
            content = r.get("content", "")[:400]
            context_parts.append(f"[{i}] {title} ({url})\n{content}")
        context_block = "\n\n".join(context_parts)

        # Truncate to max chars
        if len(context_block) > WEBSEARCH_CONTEXT_MAX_CHARS:
            context_block = context_block[:WEBSEARCH_CONTEXT_MAX_CHARS] + "…"

        enriched_prompt = f"{context_block}\n\n---\n{query_str}"

        turns = self._context_store.get_turns(user_id)
        model = self._get_user_model(user_id)
        system_instruction = self._agent_system_instruction("chat", locale)

        try:
            ollama_response = await self._generate_response(
                user_id=user_id,
                model=model,
                prompt=enriched_prompt,
                turns=turns,
                system_instruction=system_instruction,
            )
        except OllamaTimeoutError:
            await update.effective_message.reply_text(
                self._warning(self._i18n.t("errors.ollama_timeout", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaConnectionError:
            await update.effective_message.reply_text(
                self._error(self._i18n.t("errors.ollama_connection", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaError:
            await update.effective_message.reply_text(
                self._error(self._i18n.t("errors.ollama_generic", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        self._context_store.append(user_id, role="user", content=query_str)
        self._context_store.append(user_id, role="assistant", content=ollama_response.text)

        for chunk in split_message(ollama_response.text):
            await update.effective_message.reply_text(chunk, parse_mode=ParseMode.HTML)

        # Sources footer
        sources_lines = [
            f"\n{self._i18n.t('web_search.sources_header', locale=locale)}"
        ]
        for i, r in enumerate(results, 1):
            title = r.get("title", "")
            url = r.get("url", "")
            sources_lines.append(f"[{i}] <a href=\"{url}\">{title or url}</a>")
        await update.effective_message.reply_text(
            "\n".join(sources_lines),
            parse_mode=ParseMode.HTML,
            link_preview_options=LinkPreviewOptions(is_disabled=True),
            reply_markup=self._main_keyboard(locale),
        )

    async def select_file_callback(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not await self._guard_access(update):
            return
        query = update.callback_query
        if not query or not update.effective_user or not query.message:
            return

        locale = self._locale(update)
        user_id = update.effective_user.id
        await query.answer()

        data = query.data or ""
        if not data.startswith(FILE_CALLBACK_PREFIX):
            return

        payload = data.removeprefix(FILE_CALLBACK_PREFIX).strip()
        parts = payload.split(":")
        if not parts:
            return

        action = parts[0]
        if action == FILE_CLOSE_ACTION:
            try:
                await query.message.delete()
            except Exception as error:
                logger.debug("Failed to delete files message: %s", error)
                await query.edit_message_reply_markup(reply_markup=None)
            return

        if action == FILE_PAGE_ACTION and len(parts) >= 2 and parts[1].isdigit():
            await self._show_files_page(update=update, page=int(parts[1]))
            return

        if action in {FILE_TOGGLE_ACTION, FILE_DELETE_ACTION} and len(parts) >= 3:
            asset_id_raw, page_raw = parts[1], parts[2]
            if not asset_id_raw.isdigit() or not page_raw.isdigit():
                return

            asset_id = int(asset_id_raw)
            page = int(page_raw)

            try:
                if action == FILE_TOGGLE_ACTION:
                    asset = self._user_assets_store.get_asset(user_id, asset_id)
                    if not asset:
                        await query.message.reply_text(
                            self._warning(self._i18n.t("files.not_found", locale=locale)),
                            reply_markup=self._main_keyboard(locale),
                        )
                        return
                    self._user_assets_store.set_selected(user_id, asset_id, not asset.is_selected)
                else:
                    deleted = self._user_assets_store.delete_asset(user_id, asset_id)
                    if not deleted:
                        await query.message.reply_text(
                            self._warning(self._i18n.t("files.not_found", locale=locale)),
                            reply_markup=self._main_keyboard(locale),
                        )
                        return
            except Exception as error:
                logger.exception("Failed to update file action=%s asset_id=%s error=%s", action, asset_id, error)
                await query.message.reply_text(
                    self._error(self._i18n.t("errors.files_storage", locale=locale)),
                    reply_markup=self._main_keyboard(locale),
                )
                return

            await self._show_files_page(update=update, page=page)
            return

        if action == FILE_ASK_ACTION and len(parts) >= 2:
            asset_id_raw = parts[1]
            if not asset_id_raw.isdigit():
                return
            asset_id = int(asset_id_raw)
            try:
                asset = self._user_assets_store.get_asset(user_id, asset_id)
            except Exception as error:
                logger.exception("Failed to fetch file for ask action: %s", error)
                await query.message.reply_text(
                    self._error(self._i18n.t("errors.files_storage", locale=locale)),
                    reply_markup=self._main_keyboard(locale),
                )
                return

            if not asset:
                await query.message.reply_text(
                    self._warning(self._i18n.t("files.not_found", locale=locale)),
                    reply_markup=self._main_keyboard(locale),
                )
                return

            self._askfile_target_by_user[user_id] = asset_id
            await query.message.reply_text(
                self._info(
                    self._i18n.t(
                        "messages.askfile_prompt",
                        locale=locale,
                        id=asset.id,
                        name=asset.asset_name,
                    )
                ),
                reply_markup=self._main_keyboard(locale),
            )
            return

        if action == FILE_UPLOAD_ACTION:
            self._upload_mode_users.add(user_id)
            await query.message.reply_text(
                self._info(self._i18n.t("files.upload_prompt", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        if action == FILE_PREVIEW_ACTION and len(parts) >= 2:
            asset_id_raw = parts[1]
            if not asset_id_raw.isdigit():
                return
            asset_id = int(asset_id_raw)
            try:
                asset = self._user_assets_store.get_asset(user_id, asset_id)
            except Exception as error:
                logger.exception("Failed to fetch file for preview: %s", error)
                await query.message.reply_text(
                    self._error(self._i18n.t("errors.files_storage", locale=locale)),
                    reply_markup=self._main_keyboard(locale),
                )
                return
            if not asset or not asset.image_base64:
                await query.message.reply_text(
                    self._warning(self._i18n.t("files.not_found", locale=locale)),
                    reply_markup=self._main_keyboard(locale),
                )
                return
            try:
                img_bytes = base64.b64decode(asset.image_base64)
                await query.message.reply_photo(photo=img_bytes)
            except Exception as error:
                logger.warning(
                    "image_preview_failed user_id=%s asset_id=%s error=%s", user_id, asset_id, error
                )
                await query.message.reply_text(
                    self._error(self._i18n.t("errors.ollama_generic", locale=locale)),
                    reply_markup=self._main_keyboard(locale),
                )
            return

    async def _show_files_page(self, update: Update, page: int) -> None:
        query = update.callback_query
        if not query or not update.effective_user or not query.message:
            return

        locale = self._locale(update)
        user_id = update.effective_user.id
        try:
            assets = self._user_assets_store.list_assets(user_id)
        except Exception as error:
            logger.exception("Failed to list files for pagination: %s", error)
            await query.message.reply_text(
                self._error(self._i18n.t("errors.files_storage", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        if not assets:
            await self._edit_models_message(
                query=query,
                text=self._info(self._i18n.t("files.empty", locale=locale)),
                reply_markup=InlineKeyboardMarkup(
                    [
                        [
                            InlineKeyboardButton(
                                text=self._i18n.t("ui.buttons.add_file", locale=locale),
                                callback_data=f"{FILE_CALLBACK_PREFIX}{FILE_UPLOAD_ACTION}",
                            )
                        ],
                        [
                            InlineKeyboardButton(
                                text=self._i18n.t("ui.buttons.close", locale=locale),
                                callback_data=f"{FILE_CALLBACK_PREFIX}{FILE_CLOSE_ACTION}",
                            )
                        ],
                    ]
                ),
            )
            return

        page_assets, total_pages = self._paginate_assets(assets, page)
        safe_page = min(max(page, 1), total_pages)
        text = self._files_page_text(
            locale=locale,
            assets=page_assets,
            page=safe_page,
            total_pages=total_pages,
        )
        await self._edit_models_message(
            query=query,
            text=text,
            reply_markup=self._files_inline_keyboard(
                locale=locale,
                assets=page_assets,
                page=safe_page,
                total_pages=total_pages,
            ),
        )

    async def _show_web_models_page(self, update: Update, page: int, force_refresh: bool = False) -> None:
        query = update.callback_query
        if not query or not update.effective_user or not query.message:
            return

        locale = self._locale(update)
        user_id = update.effective_user.id

        try:
            models = await self._fetch_web_models(force_refresh=force_refresh)
        except OllamaTimeoutError:
            await query.message.reply_text(
                self._warning(self._i18n.t("errors.ollama_timeout", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaConnectionError:
            await query.message.reply_text(
                self._error(self._i18n.t("errors.ollama_connection", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaError as error:
            logger.warning("Ollama error while paginating web models: %s", error)
            await query.message.reply_text(
                self._error(self._i18n.t("errors.ollama_list_web_models", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        search_query = self._web_model_search_query_by_user.get(user_id, "")
        filtered_models = self._filter_web_models(models, search_query)
        if not filtered_models:
            await query.message.reply_text(
                self._warning(self._i18n.t("web_models.no_matches", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        page_models, total_pages = self._paginate_items(filtered_models, page, WEB_MODELS_PAGE_SIZE)
        safe_page = min(max(page, 1), total_pages)

        lines = [self._info(self._i18n.t("web_models.available_title", locale=locale))]
        for m in page_models:
            badges = []
            if "vision" in m.capabilities:
                badges.append("👁")
            if "thinking" in m.capabilities:
                badges.append("💭")
            line = f"- {m.name}"
            if badges:
                line += "  " + " ".join(badges)
            if m.sizes:
                line += "  📦 " + " · ".join(m.sizes[:4])
            lines.append(line)
        lines.append("")
        lines.append(self._i18n.t("web_models.page_status", locale=locale, page=safe_page, pages=total_pages))
        lines.append(self._i18n.t("web_models.select_with", locale=locale))

        await self._edit_models_message(
            query=query,
            text="\n".join(lines),
            reply_markup=self._web_models_inline_keyboard(
                locale=locale,
                models=page_models,
                page=safe_page,
                total_pages=total_pages,
            ),
        )

    async def _reply_web_models_page(self, update: Update, page: int) -> None:
        """Send (not edit) a fresh web models page — used from text handlers."""
        if not update.effective_message or not update.effective_user:
            return
        locale = self._locale(update)
        user_id = update.effective_user.id

        try:
            models = await self._fetch_web_models()
        except OllamaTimeoutError:
            await update.effective_message.reply_text(
                self._warning(self._i18n.t("errors.ollama_timeout", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaConnectionError:
            await update.effective_message.reply_text(
                self._error(self._i18n.t("errors.ollama_connection", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaError as error:
            logger.warning("Ollama error in reply_web_models_page: %s", error)
            await update.effective_message.reply_text(
                self._error(self._i18n.t("errors.ollama_list_web_models", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        search_query = self._web_model_search_query_by_user.get(user_id, "")
        filtered_models = self._filter_web_models(models, search_query)
        if not filtered_models:
            await update.effective_message.reply_text(
                self._warning(self._i18n.t("web_models.no_matches", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        page_models, total_pages = self._paginate_items(filtered_models, page, WEB_MODELS_PAGE_SIZE)
        safe_page = min(max(page, 1), total_pages)

        lines = [self._info(self._i18n.t("web_models.available_title", locale=locale))]
        for m in page_models:
            badges = []
            if "vision" in m.capabilities:
                badges.append("👁")
            if "thinking" in m.capabilities:
                badges.append("💭")
            line = f"- {m.name}"
            if badges:
                line += "  " + " ".join(badges)
            if m.sizes:
                line += "  📦 " + " · ".join(m.sizes[:4])
            lines.append(line)
        lines.append("")
        lines.append(
            self._i18n.t("web_models.page_status", locale=locale, page=safe_page, pages=total_pages)
        )
        lines.append(self._i18n.t("web_models.select_with", locale=locale))

        await update.effective_message.reply_text(
            "\n".join(lines),
            reply_markup=self._web_models_inline_keyboard(
                locale=locale,
                models=page_models,
                page=safe_page,
                total_pages=total_pages,
            ),
        )

    async def _fetch_web_models(self, force_refresh: bool = False) -> list[WebModelInfo]:
        """Return cached web model list, refreshing if the TTL has expired."""
        now = monotonic()
        if not force_refresh and self._web_models_cache and now < self._web_models_cache_expires:
            return self._web_models_cache
        models = await self._ollama_client.list_web_models()
        self._web_models_cache = models
        self._web_models_cache_expires = now + _WEB_MODELS_CACHE_TTL
        for model in models:
            self._web_model_token(model.name)
        return models

    def _web_model_token(self, model_name: str) -> str:
        existing = self._web_model_name_to_token.get(model_name)
        if existing:
            return existing

        digest = hashlib.sha1(model_name.encode("utf-8")).hexdigest()
        token_len = 10
        token = digest[:token_len]
        while self._web_model_token_to_name.get(token, model_name) != model_name:
            token_len += 1
            token = digest[:token_len]

        self._web_model_name_to_token[model_name] = token
        self._web_model_token_to_name[token] = model_name
        return token

    def _resolve_web_model_callback_value(self, value: str) -> str:
        if not value:
            return ""
        return self._web_model_token_to_name.get(value, value)

    def _filter_web_models(self, models: list[WebModelInfo], query: str) -> list[WebModelInfo]:
        search = query.strip().lower()
        if not search:
            return models
        return [
            m for m in models
            if search in m.name.lower()
            or search in m.description.lower()
            or any(search in cap for cap in m.capabilities)
            or any(search in size for size in m.sizes)
        ]

    @staticmethod
    def _format_web_model_detail(info: WebModelInfo | None, model_name: str) -> str:
        _CAP_ICONS = {
            "vision": "\U0001F441",
            "tools": "\U0001F527",
            "thinking": "\U0001F4AD",
            "embedding": "\U0001F4CA",
            "cloud": "\u2601\ufe0f",
        }
        if info is None:
            return f"\u2139\ufe0f {model_name}"
        lines: list[str] = [f"\u2139\ufe0f {info.name}"]
        if info.description:
            lines.append(info.description)
        lines.append("")
        if info.capabilities:
            caps = "  \u00b7  ".join(
                f"{_CAP_ICONS.get(c, '')} {c}" for c in info.capabilities
            )
            lines.append(caps)
        if info.sizes:
            lines.append("\U0001F4E6  " + " \u00b7 ".join(info.sizes))
        meta: list[str] = []
        if info.pulls:
            meta.append(f"\u2b07\ufe0f {info.pulls} pulls")
        if info.tags_count:
            meta.append(f"{info.tags_count} tags")
        if meta:
            lines.append("  \u00b7  ".join(meta))
        if info.updated:
            lines.append(f"\U0001F550 Updated {info.updated}")
        return "\n".join(lines)

    async def _show_models_page(self, update: Update, page: int) -> None:
        query = update.callback_query
        if not query or not update.effective_user or not query.message:
            return

        locale = self._locale(update)
        user_id = update.effective_user.id

        try:
            models = await self._ollama_client.list_models()
        except OllamaTimeoutError:
            await query.message.reply_text(
                self._warning(self._i18n.t("errors.ollama_timeout", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaConnectionError:
            await query.message.reply_text(
                self._error(self._i18n.t("errors.ollama_connection", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaError as error:
            logger.warning("Ollama error while paginating models: %s", error)
            await query.message.reply_text(
                self._error(self._i18n.t("errors.ollama_list_models", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        search_query = self._model_search_query_by_user.get(user_id, "")
        filtered_models = self._filter_models(models, search_query)
        if not filtered_models:
            await query.message.reply_text(
                self._warning(self._i18n.t("models.no_matches", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        page_models, total_pages = self._paginate_models(filtered_models, page)
        safe_page = min(max(page, 1), total_pages)
        current_model = self._get_user_model(user_id)

        lines = [self._info(self._i18n.t("models.available_title", locale=locale))]
        for model in page_models:
            marker = self._i18n.t("models.current_marker", locale=locale) if model == current_model else ""
            lines.append(f"- {model}{marker}")
        lines.append("")
        lines.append(self._i18n.t("models.page_status", locale=locale, page=safe_page, pages=total_pages))
        lines.append(self._i18n.t("models.select_with", locale=locale))
        lines.append(self._i18n.t("models.tap_button", locale=locale))

        await self._edit_models_message(
            query=query,
            text="\n".join(lines),
            reply_markup=self._models_inline_keyboard(
                locale,
                models=page_models,
                current_model=current_model,
                page=safe_page,
                total_pages=total_pages,
            ),
        )

    async def _edit_models_message(
        self,
        *,
        query,
        text: str,
        reply_markup: InlineKeyboardMarkup,
    ) -> None:
        try:
            await query.edit_message_text(
                text=text,
                reply_markup=reply_markup,
            )
        except Exception as error:
            logger.debug("Failed to edit models message, sending new message instead: %s", error)
            if query.message:
                await query.message.reply_text(
                    text,
                    reply_markup=reply_markup,
                )

    async def current_model(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not await self._guard_access(update):
            return
        if not update.effective_message or not update.effective_user:
            return

        locale = self._locale(update)
        self._log_user_event("command_currentmodel", update)
        user_id = update.effective_user.id
        current_model = self._get_user_model(user_id)
        await update.effective_message.reply_text(
            self._info(self._i18n.t("messages.current_model", locale=locale, model=current_model)),
            reply_markup=InlineKeyboardMarkup(
                [
                    [
                        InlineKeyboardButton(
                            text=self._i18n.t("ui.buttons.open_models", locale=locale),
                            callback_data=f"{MODEL_CALLBACK_PREFIX}{MODEL_REFRESH_ACTION}",
                        ),
                        InlineKeyboardButton(
                            text=self._i18n.t("ui.buttons.use_default", locale=locale),
                            callback_data=f"{MODEL_CALLBACK_PREFIX}{MODEL_DEFAULT_ACTION}",
                        ),
                    ]
                ]
            ),
        )

    async def quick_actions(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not await self._guard_access(update):
            return
        if not update.effective_message:
            return

        text = (update.effective_message.text or "").strip()
        action = self._quick_action_map.get(text)
        if action == "models":
            await self.models(update, context)
            return
        if action == "web_models":
            await self.web_models(update, context)
            return
        if action == "files":
            await self.files(update, context)
            return
        if action == "current_model":
            await self.current_model(update, context)
            return
        if action == "clear":
            await self.clear(update, context)
            return
        if action == "help":
            await self.help(update, context)
            return

    async def on_text(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not await self._guard_access(update, apply_rate_limit=True):
            return
        if not update.effective_message or not update.effective_user:
            return

        locale = self._locale(update)
        user_text = (update.effective_message.text or "").strip()
        self._log_user_event("message_received", update)
        if not user_text:
            await update.effective_message.reply_text(
                self._warning(self._i18n.t("messages.please_send_non_empty", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        user_id = update.effective_user.id
        pending_asset_id = self._askfile_target_by_user.pop(user_id, None)
        if pending_asset_id is not None:
            try:
                pending_asset = self._user_assets_store.get_asset(user_id, pending_asset_id)
            except Exception as error:
                logger.exception("Failed to load pending askfile asset: %s", error)
                await update.effective_message.reply_text(
                    self._error(self._i18n.t("errors.files_storage", locale=locale)),
                    reply_markup=self._main_keyboard(locale),
                )
                return

            if not pending_asset:
                await update.effective_message.reply_text(
                    self._warning(self._i18n.t("files.not_found", locale=locale)),
                    reply_markup=self._main_keyboard(locale),
                )
                return

            await self._answer_with_asset(
                update=update,
                locale=locale,
                user_id=user_id,
                asset=pending_asset,
                prompt=user_text,
            )
            return

        # Web model search mode: next text message is a search query for /webmodels
        if user_id in self._web_model_search_mode_users:
            self._web_model_search_mode_users.discard(user_id)
            self._web_model_search_query_by_user[user_id] = user_text
            await self._reply_web_models_page(update, 1)
            return

        turns = self._context_store.get_turns(user_id)
        model = self._get_user_model(user_id)
        started_at = monotonic()
        agent_name = self._select_agent(user_text)
        system_instruction = self._agent_system_instruction(agent_name, locale)

        # Retrieve selected assets and split by kind
        try:
            selected_assets = self._user_assets_store.search_selected_assets(
                user_id=user_id,
                query=user_text,
                limit=self._files_context_max_items,
                max_chars_total=self._files_context_max_chars,
            )
        except Exception as err:
            logger.warning("selected_assets_context_failed user_id=%s error=%s", user_id, err)
            selected_assets = []

        image_assets = [a for a in selected_assets if a.asset_kind == "image"]
        doc_assets = [a for a in selected_assets if a.asset_kind != "image"]

        # Image assets: text descriptions as prior history turns (context for all models)
        # + real image bytes attached to the current message (vision models see them directly)
        extra_turns = self._build_image_context_turns(image_assets) or None
        stored_images = [a.image_base64 for a in image_assets if a.image_base64]
        prompt_images: list[str] | None = stored_images or None
        orch_notification: str | None = None

        # Orchestrate model selection: vision when images are attached, code/general otherwise
        model, orch_notification, vision_found = await self._orchestrate_model(
            prompt=user_text,
            has_images=bool(prompt_images),
            preferred_model=model,
            locale=locale,
        )
        if prompt_images and not vision_found:
            # No vision model available — fall back to text descriptions (already in extra_turns)
            prompt_images = None
            orch_notification = None
            logger.warning(
                "ask_files no_vision_model_found model=%s dropping images, using text descriptions",
                model,
            )

        # Doc assets augment the prompt text
        if doc_assets:
            prompt_to_send = self._augment_prompt_with_assets(
                prompt=user_text, assets=doc_assets, force_single=False
            )
        else:
            prompt_to_send = user_text

        await update.effective_chat.send_action(action=ChatAction.TYPING)

        try:
            ollama_response = await self._generate_response(
                user_id=user_id,
                model=model,
                prompt=prompt_to_send,
                turns=turns,
                system_instruction=system_instruction,
                extra_turns=extra_turns,
                prompt_images=prompt_images,
            )
        except OllamaTimeoutError:
            await update.effective_message.reply_text(
                self._warning(self._i18n.t("errors.ollama_timeout", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaConnectionError:
            await update.effective_message.reply_text(
                self._error(self._i18n.t("errors.ollama_connection", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaError as error:
            logger.warning("Ollama error: %s", error)
            await update.effective_message.reply_text(
                self._error(self._i18n.t("errors.ollama_generic", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        self._context_store.append(user_id, role="user", content=user_text)
        self._context_store.append(user_id, role="assistant", content=ollama_response.text)

        elapsed_ms = int((monotonic() - started_at) * 1000)
        logger.info(
            "message_completed user_id=%s model=%s agent=%s input_chars=%d output_chars=%d elapsed_ms=%d",
            user_id,
            model,
            agent_name,
            len(user_text),
            len(ollama_response.text),
            elapsed_ms,
        )

        for chunk in split_message(ollama_response.text):
            await update.effective_message.reply_text(chunk, parse_mode=ParseMode.HTML)
        if orch_notification:
            await update.effective_message.reply_text(
                f"<i>{orch_notification}</i>", parse_mode=ParseMode.HTML
            )

    async def on_image(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not await self._guard_access(update, apply_rate_limit=True):
            return
        if not update.effective_message or not update.effective_user:
            return

        locale = self._locale(update)
        message = update.effective_message
        user_id = update.effective_user.id

        # --- UPLOAD MODE: save image to /files without model analysis ---
        if user_id in self._upload_mode_users:
            self._upload_mode_users.discard(user_id)
            try:
                upload_bytes: bytes | None = None
                upload_name = "telegram-photo"
                upload_mime = "image/jpeg"
                if message.photo:
                    sz = int(message.photo[-1].file_size or 0)
                    if sz > self._image_max_bytes:
                        await message.reply_text(
                            self._warning(
                                self._i18n.t(
                                    "image.too_large",
                                    locale=locale,
                                    max_size=self._format_size(self._image_max_bytes),
                                )
                            ),
                            reply_markup=self._main_keyboard(locale),
                        )
                        return
                    f = await message.photo[-1].get_file()
                    upload_bytes = bytes(await f.download_as_bytearray())
                elif message.document and (message.document.mime_type or "").startswith("image/"):
                    sz = int(message.document.file_size or 0)
                    if sz > self._image_max_bytes:
                        await message.reply_text(
                            self._warning(
                                self._i18n.t(
                                    "image.too_large",
                                    locale=locale,
                                    max_size=self._format_size(self._image_max_bytes),
                                )
                            ),
                            reply_markup=self._main_keyboard(locale),
                        )
                        return
                    f = await message.document.get_file()
                    upload_bytes = bytes(await f.download_as_bytearray())
                    upload_name = message.document.file_name or upload_name
                    upload_mime = message.document.mime_type or upload_mime
                if not upload_bytes:
                    await message.reply_text(
                        self._warning(self._i18n.t("image.invalid_file", locale=locale)),
                        reply_markup=self._main_keyboard(locale),
                    )
                    return
                image_b64_upload = base64.b64encode(upload_bytes).decode("utf-8")
                asset_id = self._user_assets_store.add_asset(
                    user_id=user_id,
                    asset_kind="image",
                    asset_name=upload_name,
                    mime_type=upload_mime,
                    size_bytes=len(upload_bytes),
                    content_text="",
                    is_selected=True,
                    image_base64=image_b64_upload,
                )
                await message.reply_text(
                    self._success(
                        self._i18n.t("files.upload_done", locale=locale, name=upload_name, id=asset_id)
                    ),
                    reply_markup=self._main_keyboard(locale),
                )
            except Exception as err:
                logger.warning("image_upload_save_failed user_id=%s error=%s", user_id, err)
                await message.reply_text(
                    self._error(self._i18n.t("errors.ollama_generic", locale=locale)),
                    reply_markup=self._main_keyboard(locale),
                )
            return
        # --- END UPLOAD MODE ---

        caption = (message.caption or "").strip()
        user_prompt = caption or self._i18n.t("image.default_prompt", locale=locale)
        user_prompt_with_assets = self._augment_prompt_with_selected_assets(
            user_id=user_id,
            prompt=user_prompt,
            asset_kinds={"document"},
        )
        model = self._get_user_model(user_id)
        turns = self._context_store.get_turns(user_id)
        agent_name = self._select_agent(user_prompt)
        system_instruction = self._agent_system_instruction(agent_name, locale)
        turns_for_model: list[ConversationTurn] = [
            ConversationTurn(role="system", content=system_instruction),
            *turns,
        ]
        image_bytes_size = 0

        logger.info(
            "on_image user_id=%s model=%s use_chat_api=%s locale=%s caption_len=%d has_photo=%s has_doc=%s",
            user_id, model, self._use_chat_api, locale, len(caption),
            bool(message.photo), bool(message.document),
        )

        # Orchestrate: auto-select vision-capable model or warn user
        model, orch_notification, vision_found = await self._orchestrate_model(
            prompt=user_prompt,
            has_images=True,
            preferred_model=model,
            locale=locale,
        )
        if not vision_found:
            await message.reply_text(
                self._warning(self._i18n.t("image.model_without_vision", locale=locale, model=model)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        try:
            photo_bytes: bytes | None = None
            if message.photo:
                image_bytes_size = int(message.photo[-1].file_size or 0)
                if image_bytes_size > self._image_max_bytes:
                    await message.reply_text(
                        self._warning(
                            self._i18n.t(
                                "image.too_large",
                                locale=locale,
                                max_size=self._format_size(self._image_max_bytes),
                            )
                        ),
                        reply_markup=self._main_keyboard(locale),
                    )
                    return
                file = await message.photo[-1].get_file()
                photo_bytes = bytes(await file.download_as_bytearray())
            elif message.document and (message.document.mime_type or "").startswith("image/"):
                image_bytes_size = int(message.document.file_size or 0)
                if image_bytes_size > self._image_max_bytes:
                    await message.reply_text(
                        self._warning(
                            self._i18n.t(
                                "image.too_large",
                                locale=locale,
                                max_size=self._format_size(self._image_max_bytes),
                            )
                        ),
                        reply_markup=self._main_keyboard(locale),
                    )
                    return
                file = await message.document.get_file()
                photo_bytes = bytes(await file.download_as_bytearray())

            if not photo_bytes:
                await message.reply_text(
                    self._warning(self._i18n.t("image.invalid_file", locale=locale)),
                    reply_markup=self._main_keyboard(locale),
                )
                return

            image_base64 = base64.b64encode(photo_bytes).decode("utf-8")
            if not image_bytes_size:
                image_bytes_size = len(photo_bytes)
            logger.info(
                "on_image encoded user_id=%s raw_bytes=%d b64_chars=%d",
                user_id, image_bytes_size, len(image_base64),
            )
            if image_bytes_size > self._image_max_bytes:
                await message.reply_text(
                    self._warning(
                        self._i18n.t(
                            "image.too_large",
                            locale=locale,
                            max_size=self._format_size(self._image_max_bytes),
                        )
                    ),
                    reply_markup=self._main_keyboard(locale),
                )
                return

            started_at = monotonic()
            await update.effective_chat.send_action(action=ChatAction.TYPING)

            logger.info(
                "on_image calling chat_with_image user_id=%s model=%s prompt_chars=%d context_turns=%d",
                user_id, model, len(user_prompt_with_assets), len(turns_for_model),
            )
            ollama_response = await self._ollama_client.chat_with_image(
                model=model,
                prompt=user_prompt_with_assets,
                images=[image_base64],
                context_turns=turns_for_model,
                keep_alive=self._keep_alive,
            )
            logger.info(
                "on_image response user_id=%s model=%s response_chars=%d",
                user_id, model, len(ollama_response.text),
            )
        except OllamaTimeoutError:
            await message.reply_text(
                self._warning(self._i18n.t("errors.ollama_timeout", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaConnectionError:
            await message.reply_text(
                self._error(self._i18n.t("errors.ollama_connection", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except OllamaError as error:
            logger.warning("Ollama image error: %s", error)
            await message.reply_text(
                self._error(self._i18n.t("image.processing_error", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return
        except Exception as error:
            logger.warning("image_read_failed user_id=%s error=%s", user_id, error)
            await message.reply_text(
                self._warning(self._i18n.t("image.read_error", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
            return

        self._context_store.append(user_id, role="user", content=f"[Image] {user_prompt}")
        self._context_store.append(user_id, role="assistant", content=ollama_response.text)
        try:
            self._user_assets_store.add_asset(
                user_id=user_id,
                asset_kind="image",
                asset_name=(message.document.file_name if message.document else "telegram-photo"),
                mime_type=(message.document.mime_type if message.document else "image/jpeg") or "image/jpeg",
                size_bytes=image_bytes_size,
                content_text=ollama_response.text,
                is_selected=True,
                image_base64=image_base64,
            )
        except Exception as error:
            logger.warning("image_asset_save_failed user_id=%s error=%s", user_id, error)

        elapsed_ms = int((monotonic() - started_at) * 1000)
        logger.info(
            "image_completed user_id=%s model=%s agent=%s prompt_chars=%d output_chars=%d elapsed_ms=%d",
            user_id,
            model,
            agent_name,
            len(user_prompt),
            len(ollama_response.text),
            elapsed_ms,
        )

        for chunk in split_message(ollama_response.text):
            await message.reply_text(chunk, parse_mode=ParseMode.HTML)
        if orch_notification:
            await message.reply_text(f"<i>{orch_notification}</i>", parse_mode=ParseMode.HTML)

    async def on_document(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not await self._guard_access(update, apply_rate_limit=True):
            return
        if not update.effective_message or not update.effective_user:
            return

        locale = self._locale(update)
        message = update.effective_message
        document = message.document
        if not document:
            return

        user_id = update.effective_user.id
        file_name = document.file_name or "document"
        file_size = int(document.file_size or 0)
        mime_type = (document.mime_type or "").lower()

        if file_size > self._document_max_bytes:
            await message.reply_text(
                self._warning(
                    self._i18n.t(
                        "document.too_large",
                        locale=locale,
                        max_size=self._format_size(self._document_max_bytes),
                    )
                ),
                reply_markup=self._main_keyboard(locale),
            )
            return

        try:
            telegram_file = await document.get_file()
            raw_bytes = bytes(await telegram_file.download_as_bytearray())
            if len(raw_bytes) > self._document_max_bytes:
                await message.reply_text(
                    self._warning(
                        self._i18n.t(
                            "document.too_large",
                            locale=locale,
                            max_size=self._format_size(self._document_max_bytes),
                        )
                    ),
                    reply_markup=self._main_keyboard(locale),
                )
                return

            extracted_text = self._extract_document_text(
                content=raw_bytes,
                mime_type=mime_type,
                file_name=file_name,
            )
            if not extracted_text.strip():
                await message.reply_text(
                    self._warning(self._i18n.t("document.empty", locale=locale)),
                    reply_markup=self._main_keyboard(locale),
                )
                return

            trimmed_text = self._trim_document_text(extracted_text)

            # In upload mode, ignore the caption and save without model analysis
            if user_id in self._upload_mode_users:
                self._upload_mode_users.discard(user_id)
                caption = ""
            else:
                caption = (message.caption or "").strip()

            asset_id: int | None = None
            try:
                asset_id = self._user_assets_store.add_asset(
                    user_id=user_id,
                    asset_kind="document",
                    asset_name=file_name,
                    mime_type=mime_type or "application/octet-stream",
                    size_bytes=len(raw_bytes),
                    content_text=trimmed_text,
                    is_selected=True,
                )
            except Exception as error:
                logger.warning("document_asset_save_failed user_id=%s file_name=%s error=%s", user_id, file_name, error)

            if not caption:
                self._context_store.append(
                    user_id,
                    role="user",
                    content=f"[Document: {file_name}]\n{trimmed_text}",
                )
                await message.reply_text(
                    self._success(self._i18n.t("document.added", locale=locale, name=file_name, id=asset_id or "?")),
                    reply_markup=self._main_keyboard(locale),
                )
                return

            turns = self._context_store.get_turns(user_id)
            model = self._get_user_model(user_id)
            agent_name = self._select_agent(caption)
            system_instruction = self._agent_system_instruction(agent_name, locale)
            started_at = monotonic()

            review_prompt = (
                f"Document name: {file_name}\n\n"
                f"Document content:\n{trimmed_text}\n\n"
                f"User request: {caption}"
            )

            selected_model, orch_notification, _ = await self._orchestrate_model(
                prompt=review_prompt,
                has_images=False,
                preferred_model=model,
                locale=locale,
            )
            model = selected_model

            await update.effective_chat.send_action(action=ChatAction.TYPING)
            ollama_response = await self._generate_response(
                user_id=user_id,
                model=model,
                prompt=review_prompt,
                turns=turns,
                system_instruction=system_instruction,
            )

            self._context_store.append(
                user_id,
                role="user",
                content=f"[Document review: {file_name}] {caption}",
            )
            self._context_store.append(user_id, role="assistant", content=ollama_response.text)

            elapsed_ms = int((monotonic() - started_at) * 1000)
            logger.info(
                "document_review_completed user_id=%s model=%s file_name=%s input_chars=%d output_chars=%d elapsed_ms=%d",
                user_id,
                model,
                file_name,
                len(review_prompt),
                len(ollama_response.text),
                elapsed_ms,
            )

            for chunk in split_message(ollama_response.text):
                await message.reply_text(chunk, parse_mode=ParseMode.HTML)
            if orch_notification:
                await message.reply_text(f"<i>{orch_notification}</i>", parse_mode=ParseMode.HTML)
        except ValueError as error:
            logger.warning("document_unsupported user_id=%s file_name=%s error=%s", user_id, file_name, error)
            await message.reply_text(
                self._warning(self._i18n.t("document.unsupported", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
        except OllamaTimeoutError:
            await message.reply_text(
                self._warning(self._i18n.t("errors.ollama_timeout", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
        except OllamaConnectionError:
            await message.reply_text(
                self._error(self._i18n.t("errors.ollama_connection", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
        except OllamaError as error:
            logger.warning("document_ollama_error user_id=%s file_name=%s error=%s", user_id, file_name, error)
            await message.reply_text(
                self._error(self._i18n.t("document.processing_error", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )
        except Exception as error:
            logger.warning("document_processing_failed user_id=%s file_name=%s error=%s", user_id, file_name, error)
            await message.reply_text(
                self._error(self._i18n.t("document.processing_error", locale=locale)),
                reply_markup=self._main_keyboard(locale),
            )

    async def on_voice_or_audio(self, update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if not await self._guard_access(update):
            return
        if not update.effective_message:
            return
        locale = self._locale(update)

        await update.effective_message.reply_text(
            self._warning(self._i18n.t("messages.voice_disabled", locale=locale)),
            reply_markup=self._main_keyboard(locale),
        )

    async def _generate_response(
        self,
        *,
        user_id: int,
        model: str,
        prompt: str,
        turns: list[ConversationTurn],
        system_instruction: str,
        extra_turns: list[ConversationTurn] | None = None,
        prompt_images: list[str] | None = None,
    ):
        """Call Ollama and return a response.

        ``extra_turns`` are injected before the real history and carry text-based
        context (e.g. stored image descriptions, document summaries).
        ``prompt_images`` are attached to the *current* user message as base64
        strings — the only placement that vision models actually process.
        """
        turns_for_model: list[ConversationTurn] = [ConversationTurn(role="system", content=system_instruction)]
        if extra_turns:
            turns_for_model.extend(extra_turns)
        turns_for_model.extend(turns)

        logger.info(
            "_generate_response user_id=%s model=%s use_chat_api=%s prompt_images=%s prompt_chars=%d context_turns=%d",
            user_id, model, self._use_chat_api,
            len(prompt_images) if prompt_images else 0,
            len(prompt), len(turns_for_model),
        )

        if self._use_chat_api:
            if prompt_images:
                # Route through chat_with_image which has _looks_like_missing_image_response
                # fallback detection — identical path to direct image upload.
                logger.info("_generate_response -> chat_with_image user_id=%s model=%s images=%d", user_id, model, len(prompt_images))
                try:
                    return await self._ollama_client.chat_with_image(
                        model=model,
                        prompt=prompt,
                        images=prompt_images,
                        context_turns=turns_for_model,
                        keep_alive=self._keep_alive,
                    )
                except OllamaError as error:
                    logger.warning(
                        "ollama_chat_image_fallback_to_generate user_id=%s model=%s error=%s",
                        user_id,
                        model,
                        error,
                    )
            else:
                logger.info("_generate_response -> chat (text-only) user_id=%s model=%s", user_id, model)
                try:
                    return await self._ollama_client.chat(
                        model=model,
                        prompt=prompt,
                        context_turns=turns_for_model,
                        keep_alive=self._keep_alive,
                    )
                except OllamaError as error:
                    logger.warning(
                        "ollama_chat_fallback_to_generate user_id=%s model=%s error=%s",
                        user_id,
                        model,
                        error,
                    )

        logger.info(
            "_generate_response -> generate (fallback) user_id=%s model=%s images=%s",
            user_id, model, len(prompt_images) if prompt_images else 0,
        )
        return await self._ollama_client.generate(
            model=model,
            prompt=prompt,
            context_turns=turns_for_model,
            images=prompt_images if prompt_images else None,
            keep_alive=self._keep_alive,
        )

    async def _orchestrate_model(
        self,
        *,
        prompt: str,
        has_images: bool,
        preferred_model: str,
        locale: str,
    ) -> tuple[str, str | None, bool]:
        """Delegate task detection and model selection to the orchestrator.

        Returns
        -------
        selected_model:
            The model that should handle this request.
        notification:
            An HTML string to send to the user when the model was changed,
            ``None`` if the preferred model is used as-is.
        found_suitable:
            ``False`` only when task=vision and no vision model is available;
            the caller should warn the user.
        """
        from src.services.model_orchestrator import TASK_VISION, TASK_CODE

        task = self._model_orchestrator.detect_task(prompt, has_images)
        selected_model, changed, found_suitable = await self._model_orchestrator.select_model(
            task, preferred_model
        )
        notification: str | None = None
        if changed:
            task_label = self._i18n.t(f"orchestrator.task_{task}", locale=locale)
            notification = self._i18n.t(
                "orchestrator.switched_model",
                locale=locale,
                model=selected_model,
                task=task_label,
            )
        logger.info(
            "orchestrate_model task=%s preferred=%s selected=%s changed=%s found_suitable=%s",
            task,
            preferred_model,
            selected_model,
            changed,
            found_suitable,
        )
        return selected_model, notification, found_suitable

    @staticmethod
    def _select_agent(user_text: str) -> str:
        text = user_text.lower()

        planning_keywords = [
            "plan",
            "roadmap",
            "step by step",
            "paso a paso",
            "planifica",
            "estrategia",
        ]
        analysis_keywords = [
            "resume",
            "resumen",
            "analiza",
            "análisis",
            "extrae",
            "clasifica",
            "category",
            "sentiment",
        ]

        if any(keyword in text for keyword in planning_keywords):
            return "planner"
        if any(keyword in text for keyword in analysis_keywords):
            return "analyst"
        return "chat"

    def _agent_system_instruction(self, agent_name: str, locale: str) -> str:
        if agent_name == "planner":
            return self._i18n.t("agent.planner_instruction", locale=locale)
        if agent_name == "analyst":
            return self._i18n.t("agent.analyst_instruction", locale=locale)
        return self._i18n.t("agent.chat_instruction", locale=locale)

    def _get_user_model(self, user_id: int) -> str:
        try:
            selected_model = self._model_preferences_store.get_user_model(user_id)
        except Exception as error:
            logger.exception("Failed to load user model preference: %s", error)
            return self._default_model

        if not selected_model:
            return self._default_model
        return selected_model

    async def _guard_access(self, update: Update, apply_rate_limit: bool = False) -> bool:
        user = update.effective_user
        if not user:
            return False

        if self._allowed_user_ids and user.id not in self._allowed_user_ids:
            logger.warning("access_denied user_id=%s", user.id)
            await self._deny_access(update)
            return False

        if apply_rate_limit and self._rate_limiter and not self._rate_limiter.allow(user.id):
            logger.warning("rate_limit_exceeded user_id=%s", user.id)
            locale = self._locale(update)
            target_message = update.effective_message or (
                update.callback_query.message if update.callback_query else None
            )
            if target_message:
                await target_message.reply_text(
                    self._warning(self._i18n.t("messages.rate_limit_exceeded", locale=locale)),
                    reply_markup=self._main_keyboard(locale),
                )
            return False

        return True

    async def _deny_access(self, update: Update) -> None:
        locale = self._locale(update)
        denied_text = self._error(self._i18n.t("messages.access_denied", locale=locale))

        query = update.callback_query
        if query:
            await query.answer(self._i18n.t("messages.access_denied_alert", locale=locale), show_alert=True)

        target_message = update.effective_message or (query.message if query else None)
        if target_message:
            await target_message.reply_text(denied_text)

    @staticmethod
    def _log_user_event(event: str, update: Update) -> None:
        user_id = update.effective_user.id if update.effective_user else "unknown"
        chat_id = update.effective_chat.id if update.effective_chat else "unknown"
        logger.info("%s user_id=%s chat_id=%s", event, user_id, chat_id)

    @staticmethod
    def _status(icon: str, text: str) -> str:
        return f"{icon} {text}"

    @staticmethod
    def _info(text: str) -> str:
        return BotHandlers._status(ICON_INFO, text)

    @staticmethod
    def _success(text: str) -> str:
        return BotHandlers._status(ICON_SUCCESS, text)

    @staticmethod
    def _warning(text: str) -> str:
        return BotHandlers._status(ICON_WARNING, text)

    @staticmethod
    def _error(text: str) -> str:
        return BotHandlers._status(ICON_ERROR, text)

    @staticmethod
    def _format_size(num_bytes: int) -> str:
        mib = num_bytes / (1024 * 1024)
        return f"{mib:.1f} MB"

    @staticmethod
    def _format_size_compact(num_bytes: int) -> str:
        if num_bytes < 1024:
            return f"{num_bytes} B"
        if num_bytes < 1024 * 1024:
            return f"{num_bytes / 1024:.0f} KB"
        return f"{num_bytes / (1024 * 1024):.1f} MB"

    @staticmethod
    def _truncate_text(text: str, *, max_chars: int) -> str:
        cleaned = text.strip()
        if len(cleaned) <= max_chars:
            return cleaned
        return f"{cleaned[: max_chars - 1]}…"

    def _asset_preview(self, asset: UserAsset) -> str:
        text = asset.content_text
        if not text.strip():
            return "—"

        if asset.asset_kind == "image":
            analysis_match = re.search(
                r"Image analysis result:\s*(.+)",
                text,
                flags=re.IGNORECASE | re.DOTALL,
            )
            if analysis_match:
                analysis_text = " ".join(analysis_match.group(1).split())
                if analysis_text:
                    return self._truncate_text(f"🖼️ {analysis_text}", max_chars=95)

            image_prompt_match = re.search(r"Image prompt:\s*(.+)", text, flags=re.IGNORECASE)
            if image_prompt_match:
                prompt = image_prompt_match.group(1).strip()
                if prompt:
                    return self._truncate_text(f"🖼️ {prompt}", max_chars=95)

            return "🖼️ Image asset"

        lines = [line.strip() for line in text.splitlines() if line.strip()]
        if lines:
            first_line = lines[0]
            if first_line.lower().startswith("document name:") and len(lines) > 1:
                first_line = lines[1]

            if len(first_line) > 8:
                return self._truncate_text(f"📄 {first_line}", max_chars=95)

        single_line = " ".join(text.split())
        return self._truncate_text(single_line, max_chars=95)

    def _clean_image_text(self, content_text: str) -> str:
        """Strip legacy format labels to return plain analysis text."""
        return re.sub(
            r"^(?:image prompt:[^\n]*\n+)?(?:image analysis result[^:\n]*:\n+)?",
            "",
            content_text.strip(),
            flags=re.IGNORECASE,
        ).strip() or content_text.strip()

    def _build_image_context_turns(self, assets: list[UserAsset]) -> list[ConversationTurn]:
        """Build synthetic user+assistant turn pairs for stored image assets.

        These turns inject the stored text analysis as prior history so that
        text-only models have descriptive context.  Vision models receive the
        actual image bytes separately via ``prompt_images`` on the *current*
        user message — Ollama only processes images in the active message, not
        in history turns, so bytes are NOT included here.
        """
        turns: list[ConversationTurn] = []
        for asset in assets:
            if asset.asset_kind != "image":
                continue
            analysis = self._clean_image_text(asset.content_text)
            if not analysis:
                continue
            turns.append(
                ConversationTurn(
                    role="user",
                    content=f"[Image uploaded: {asset.asset_name}]",
                )
            )
            turns.append(ConversationTurn(role="assistant", content=analysis))
        return turns

    def _trim_document_text(self, text: str) -> str:
        cleaned = text.strip()
        if len(cleaned) <= self._document_max_chars:
            return cleaned
        return f"{cleaned[:self._document_max_chars]}\n\n[...truncated...]"

    @staticmethod
    def _extract_document_text(*, content: bytes, mime_type: str, file_name: str) -> str:
        suffix = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
        text_extensions = {
            "txt",
            "md",
            "markdown",
            "csv",
            "json",
            "xml",
            "yml",
            "yaml",
            "log",
            "py",
            "js",
            "ts",
            "java",
            "c",
            "cpp",
            "html",
            "css",
            "sql",
            "rst",
        }

        if mime_type.startswith("text/") or suffix in text_extensions:
            return content.decode("utf-8", errors="replace")

        if mime_type == "application/pdf" or suffix == "pdf":
            try:
                from pypdf import PdfReader
            except Exception as error:
                raise ValueError("PDF extraction requires pypdf") from error

            reader = PdfReader(io.BytesIO(content))
            pages_text = [page.extract_text() or "" for page in reader.pages]
            return "\n\n".join(pages_text)

        docx_mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        if mime_type == docx_mime or suffix == "docx":
            try:
                from docx import Document as DocxDocument
            except Exception as error:
                raise ValueError("DOCX extraction requires python-docx") from error

            doc = DocxDocument(io.BytesIO(content))
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            return "\n".join(paragraphs)

        xlsx_mimes = {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel",
        }
        if mime_type in xlsx_mimes or suffix in ("xlsx", "xls"):
            try:
                import openpyxl
            except Exception as error:
                raise ValueError("XLSX extraction requires openpyxl") from error

            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            parts: list[str] = []
            for sheet in wb.worksheets:
                sheet_rows: list[str] = []
                for row in sheet.iter_rows(values_only=True):
                    row_text = "\t".join(str(cell) if cell is not None else "" for cell in row).rstrip()
                    if row_text.strip():
                        sheet_rows.append(row_text)
                if sheet_rows:
                    parts.append(f"[Sheet: {sheet.title}]\n" + "\n".join(sheet_rows))
            return "\n\n".join(parts)

        raise ValueError(f"Unsupported document format: {mime_type or suffix or 'unknown'}")

    def _locale(self, update: Update) -> str:
        language_code = update.effective_user.language_code if update.effective_user else None
        return self._i18n.resolve_locale(language_code)

    def _build_quick_action_map(self) -> dict[str, str]:
        quick_action_keys = {
            "models": "ui.buttons.models",
            "web_models": "ui.buttons.web_models",
            "files": "ui.buttons.files",
            "current_model": "ui.buttons.current_model",
            "clear": "ui.buttons.clear",
            "help": "ui.buttons.help",
        }
        mapping: dict[str, str] = {}
        for locale in self._i18n.available_locales:
            for action, key in quick_action_keys.items():
                mapping[self._i18n.t(key, locale=locale)] = action
        return mapping

    def quick_actions_regex(self) -> str:
        labels = sorted({re.escape(label) for label in self._quick_action_map})
        return rf"^({'|'.join(labels)})$"

    def _filter_models(self, models: list[str], query: str) -> list[str]:
        search = query.strip().lower()
        if not search:
            return models
        return [model for model in models if search in model.lower()]

    def _paginate_models(self, models: list[str], page: int) -> tuple[list[str], int]:
        return self._paginate_items(models, page, MODELS_PAGE_SIZE)

    def _paginate_assets(self, assets: list[UserAsset], page: int) -> tuple[list[UserAsset], int]:
        if not assets:
            return [], 1

        total_pages = max(1, (len(assets) + FILES_PAGE_SIZE - 1) // FILES_PAGE_SIZE)
        safe_page = min(max(page, 1), total_pages)
        start = (safe_page - 1) * FILES_PAGE_SIZE
        end = start + FILES_PAGE_SIZE
        return assets[start:end], total_pages

    def _files_page_text(self, *, locale: str, assets: list[UserAsset], page: int, total_pages: int) -> str:
        lines = [self._info(self._i18n.t("files.available_title", locale=locale))]
        for asset in assets:
            selected_marker = "✅" if asset.is_selected else "☑️"
            kind = "doc" if asset.asset_kind == "document" else "img"
            asset_name = self._truncate_text(asset.asset_name, max_chars=38)
            size = self._format_size_compact(asset.size_bytes)
            preview = self._asset_preview(asset)
            lines.append(f"- #{asset.id} [{kind}] {asset_name} {selected_marker} ({size})")
            lines.append(f"  {preview}")
        lines.append("")
        lines.append(self._i18n.t("files.page_status", locale=locale, page=page, pages=total_pages))
        lines.append(self._i18n.t("files.instructions", locale=locale))
        return "\n".join(lines)

    def _files_inline_keyboard(
        self,
        *,
        locale: str,
        assets: list[UserAsset],
        page: int,
        total_pages: int,
    ) -> InlineKeyboardMarkup:
        rows: list[list[InlineKeyboardButton]] = []
        for asset in assets:
            toggle_label = f"✅ #{asset.id}" if asset.is_selected else f"☑️ #{asset.id}"
            rows.append(
                [
                    InlineKeyboardButton(
                        text=f"{self._i18n.t('ui.buttons.ask_file', locale=locale)} #{asset.id}",
                        callback_data=f"{FILE_CALLBACK_PREFIX}{FILE_ASK_ACTION}:{asset.id}",
                    ),
                    InlineKeyboardButton(
                        text=toggle_label,
                        callback_data=f"{FILE_CALLBACK_PREFIX}{FILE_TOGGLE_ACTION}:{asset.id}:{page}",
                    ),
                    InlineKeyboardButton(
                        text=f"🗑 #{asset.id}",
                        callback_data=f"{FILE_CALLBACK_PREFIX}{FILE_DELETE_ACTION}:{asset.id}:{page}",
                    ),
                ]
            )
            if asset.asset_kind == "image" and asset.image_base64:
                rows.append(
                    [
                        InlineKeyboardButton(
                            text=f"{self._i18n.t('ui.buttons.preview', locale=locale)} #{asset.id}",
                            callback_data=f"{FILE_CALLBACK_PREFIX}{FILE_PREVIEW_ACTION}:{asset.id}",
                        )
                    ]
                )

        if total_pages > 1:
            nav_row: list[InlineKeyboardButton] = []
            if page > 1:
                nav_row.append(
                    InlineKeyboardButton(
                        text=self._i18n.t("ui.buttons.prev_page", locale=locale),
                        callback_data=f"{FILE_CALLBACK_PREFIX}{FILE_PAGE_ACTION}:{page - 1}",
                    )
                )
            if page < total_pages:
                nav_row.append(
                    InlineKeyboardButton(
                        text=self._i18n.t("ui.buttons.next_page", locale=locale),
                        callback_data=f"{FILE_CALLBACK_PREFIX}{FILE_PAGE_ACTION}:{page + 1}",
                    )
                )
            if nav_row:
                rows.append(nav_row)

        rows.append(
            [
                InlineKeyboardButton(
                    text=self._i18n.t("ui.buttons.add_file", locale=locale),
                    callback_data=f"{FILE_CALLBACK_PREFIX}{FILE_UPLOAD_ACTION}",
                )
            ]
        )
        rows.append(
            [
                InlineKeyboardButton(
                    text=self._i18n.t("ui.buttons.close", locale=locale),
                    callback_data=f"{FILE_CALLBACK_PREFIX}{FILE_CLOSE_ACTION}",
                )
            ]
        )

        return InlineKeyboardMarkup(rows)

    def _augment_prompt_with_selected_assets(self, *, user_id: int, prompt: str, asset_kinds: set[str] | None = None) -> str:
        try:
            selected_assets = self._user_assets_store.search_selected_assets(
                user_id=user_id,
                query=prompt,
                limit=self._files_context_max_items,
                max_chars_total=self._files_context_max_chars,
                asset_kinds=asset_kinds,
            )
        except Exception as error:
            logger.warning("selected_assets_context_failed user_id=%s error=%s", user_id, error)
            return prompt

        if not selected_assets:
            return prompt

        return self._augment_prompt_with_assets(prompt=prompt, assets=selected_assets, force_single=False)

    def _augment_prompt_with_assets(
        self,
        *,
        prompt: str,
        assets: list[UserAsset],
        force_single: bool,
    ) -> str:
        if not assets:
            return prompt

        has_image_assets = any(asset.asset_kind == "image" for asset in assets)

        lines = [
            "=== STORED FILE CONTEXT ===",
        ]
        if has_image_assets:
            lines.append(
                "The following entries include visual analyses of images previously uploaded by the user. "
                "Treat this analysis data as your direct knowledge of the image content."
            )
        if force_single:
            lines.append("Answer exclusively based on the single file provided below.")
        lines.append("")

        for asset in assets:
            if asset.asset_kind == "image":
                # Strip legacy labels from old-format content_text
                clean_text = re.sub(
                    r"^(?:image prompt:[^\n]*\n+)?(?:image analysis result[^:\n]*:\n+)?",
                    "",
                    asset.content_text.strip(),
                    flags=re.IGNORECASE,
                ).strip() or asset.content_text.strip()
                lines.append(f"[Image #{asset.id}: {asset.asset_name} — visual analysis]")
                lines.append(clean_text)
            else:
                lines.append(f"[Document #{asset.id}: {asset.asset_name}]")
                lines.append(asset.content_text.strip())
            lines.append("")

        lines.append("=== END CONTEXT ===")
        lines.append("")
        lines.append(f"User request: {prompt}")
        return "\n".join(lines)

    def _paginate_items(self, items: list[str], page: int, page_size: int) -> tuple[list[str], int]:
        if not items:
            return [], 1

        total_pages = max(1, (len(items) + page_size - 1) // page_size)
        safe_page = min(max(page, 1), total_pages)
        start = (safe_page - 1) * page_size
        end = start + page_size
        return items[start:end], total_pages

    def _models_inline_keyboard(
        self,
        locale: str,
        models: list[str],
        current_model: str,
        page: int = 1,
        total_pages: int = 1,
    ) -> InlineKeyboardMarkup:
        rows: list[list[InlineKeyboardButton]] = []
        row: list[InlineKeyboardButton] = []

        for model in models:
            label = f"✅ {model}" if model == current_model else model
            callback_data = f"{MODEL_CALLBACK_PREFIX}{model}"

            if len(callback_data) > 64:
                continue

            row.append(InlineKeyboardButton(text=label, callback_data=callback_data))
            if len(row) == 2:
                rows.append(row)
                row = []

        if row:
            rows.append(row)

        if total_pages > 1:
            nav_row: list[InlineKeyboardButton] = []
            if page > 1:
                nav_row.append(
                    InlineKeyboardButton(
                        text=self._i18n.t("ui.buttons.prev_page", locale=locale),
                        callback_data=f"{MODEL_CALLBACK_PREFIX}{MODEL_PAGE_ACTION_PREFIX}{page - 1}",
                    )
                )
            if page < total_pages:
                nav_row.append(
                    InlineKeyboardButton(
                        text=self._i18n.t("ui.buttons.next_page", locale=locale),
                        callback_data=f"{MODEL_CALLBACK_PREFIX}{MODEL_PAGE_ACTION_PREFIX}{page + 1}",
                    )
                )
            if nav_row:
                rows.append(nav_row)

        if not rows:
            rows = [
                [
                    InlineKeyboardButton(
                        text=self._i18n.t("ui.buttons.refresh_models", locale=locale),
                        callback_data=f"{MODEL_CALLBACK_PREFIX}{MODEL_REFRESH_ACTION}",
                    )
                ]
            ]

        rows.append(
            [
                InlineKeyboardButton(
                    text=self._i18n.t("ui.buttons.use_default", locale=locale),
                    callback_data=f"{MODEL_CALLBACK_PREFIX}{MODEL_DEFAULT_ACTION}",
                ),
                InlineKeyboardButton(
                    text=self._i18n.t("ui.buttons.refresh", locale=locale),
                    callback_data=f"{MODEL_CALLBACK_PREFIX}{MODEL_REFRESH_ACTION}",
                ),
                InlineKeyboardButton(
                    text=self._i18n.t("ui.buttons.close", locale=locale),
                    callback_data=f"{MODEL_CALLBACK_PREFIX}{MODEL_CLOSE_ACTION}",
                ),
            ]
        )

        return InlineKeyboardMarkup(rows)

    def _web_models_inline_keyboard(
        self,
        locale: str,
        models: list[WebModelInfo],
        page: int,
        total_pages: int,
    ) -> InlineKeyboardMarkup:
        rows: list[list[InlineKeyboardButton]] = []
        row: list[InlineKeyboardButton] = []

        for m in models:
            label = m.name
            if "vision" in m.capabilities:
                label = f"\U0001F441 {label}"
            elif "thinking" in m.capabilities:
                label = f"\U0001F4AD {label}"
            token = self._web_model_token(m.name)
            row.append(
                InlineKeyboardButton(
                    text=label,
                    callback_data=f"{WEB_MODEL_CALLBACK_PREFIX}{WEB_MODEL_DETAIL_ACTION}{token}",
                )
            )
            if len(row) == 2:
                rows.append(row)
                row = []

        if row:
            rows.append(row)

        if total_pages > 1:
            nav_row: list[InlineKeyboardButton] = []
            if page > 1:
                nav_row.append(
                    InlineKeyboardButton(
                        text=self._i18n.t("ui.buttons.prev_page", locale=locale),
                        callback_data=f"{WEB_MODEL_CALLBACK_PREFIX}{WEB_MODEL_PAGE_ACTION_PREFIX}{page - 1}",
                    )
                )
            if page < total_pages:
                nav_row.append(
                    InlineKeyboardButton(
                        text=self._i18n.t("ui.buttons.next_page", locale=locale),
                        callback_data=f"{WEB_MODEL_CALLBACK_PREFIX}{WEB_MODEL_PAGE_ACTION_PREFIX}{page + 1}",
                    )
                )
            if nav_row:
                rows.append(nav_row)

        rows.append(
            [
                InlineKeyboardButton(
                    text=self._i18n.t("ui.buttons.search", locale=locale),
                    callback_data=f"{WEB_MODEL_CALLBACK_PREFIX}{WEB_MODEL_SEARCH_ACTION}",
                ),
                InlineKeyboardButton(
                    text=self._i18n.t("ui.buttons.refresh", locale=locale),
                    callback_data=f"{WEB_MODEL_CALLBACK_PREFIX}{WEB_MODEL_REFRESH_ACTION}",
                ),
                InlineKeyboardButton(
                    text=self._i18n.t("ui.buttons.close", locale=locale),
                    callback_data=f"{WEB_MODEL_CALLBACK_PREFIX}{WEB_MODEL_CLOSE_ACTION}",
                ),
            ]
        )

        return InlineKeyboardMarkup(rows)

    def _main_keyboard(self, locale: str) -> ReplyKeyboardMarkup:
        return ReplyKeyboardMarkup(
            keyboard=[
                [
                    KeyboardButton(self._i18n.t("ui.buttons.models", locale=locale)),
                    KeyboardButton(self._i18n.t("ui.buttons.web_models", locale=locale)),
                ],
                [
                    KeyboardButton(self._i18n.t("ui.buttons.files", locale=locale)),
                    KeyboardButton(self._i18n.t("ui.buttons.current_model", locale=locale)),
                ],
                [
                    KeyboardButton(self._i18n.t("ui.buttons.clear", locale=locale)),
                    KeyboardButton(self._i18n.t("ui.buttons.help", locale=locale)),
                ],
            ],
            resize_keyboard=True,
            is_persistent=True,
            input_field_placeholder=self._i18n.t("ui.input_placeholder", locale=locale),
        )

    def _clear_inline_keyboard(self, locale: str) -> InlineKeyboardMarkup:
        return InlineKeyboardMarkup(
            [
                [
                    InlineKeyboardButton(
                        text=self._i18n.t("ui.buttons.confirm", locale=locale),
                        callback_data=f"{CLEAR_CALLBACK_PREFIX}confirm",
                    ),
                    InlineKeyboardButton(
                        text=self._i18n.t("ui.buttons.cancel", locale=locale),
                        callback_data=f"{CLEAR_CALLBACK_PREFIX}cancel",
                    ),
                ]
            ]
        )


def register_handlers(application: Application, handlers: BotHandlers) -> None:
    application.post_init = handlers.set_commands
    application.add_handler(CommandHandler("start", handlers.start))
    application.add_handler(CommandHandler("help", handlers.help))
    application.add_handler(CommandHandler("health", handlers.health))
    application.add_handler(CommandHandler("clear", handlers.clear))
    application.add_handler(CommandHandler("models", handlers.models))
    application.add_handler(CommandHandler("webmodels", handlers.web_models))
    application.add_handler(CommandHandler("files", handlers.files))
    application.add_handler(CommandHandler("askfile", handlers.askfile))
    application.add_handler(CommandHandler("cancel", handlers.cancel))
    application.add_handler(CommandHandler("currentmodel", handlers.current_model))
    application.add_handler(CommandHandler("deletemodel", handlers.delete_model))
    application.add_handler(CommandHandler("info", handlers.model_info))
    application.add_handler(CommandHandler("websearch", handlers.web_search_cmd))
    application.add_handler(CallbackQueryHandler(handlers.select_model_callback, pattern=r"^model:"))
    application.add_handler(CallbackQueryHandler(handlers.select_web_model_callback, pattern=r"^webmodel:"))
    application.add_handler(CallbackQueryHandler(handlers.select_file_callback, pattern=r"^file:"))
    application.add_handler(CallbackQueryHandler(handlers.clear_callback, pattern=r"^clear:"))
    application.add_handler(CallbackQueryHandler(handlers.delete_model_callback, pattern=r"^delmod:"))
    application.add_handler(
        MessageHandler(
            filters.Regex(handlers.quick_actions_regex()),
            handlers.quick_actions,
        )
    )
    application.add_handler(
        MessageHandler(filters.VOICE | filters.AUDIO | filters.VIDEO_NOTE, handlers.on_voice_or_audio)
    )
    application.add_handler(
        MessageHandler(filters.PHOTO | filters.Document.IMAGE, handlers.on_image)
    )
    application.add_handler(
        MessageHandler(filters.Document.ALL & ~filters.Document.IMAGE, handlers.on_document)
    )
    application.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handlers.on_text))
